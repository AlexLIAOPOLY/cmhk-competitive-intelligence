import unittest
import tempfile
from pathlib import Path
from unittest.mock import patch
from docx import Document

from tts_service import normalize_for_speech, prepare_tts_text, build_audio_summary


class PrepareTtsTextTests(unittest.TestCase):
    def test_performance_voice_reads_earnings_from_the_same_word_table(self):
        with tempfile.TemporaryDirectory() as temp:
            report = Path(temp) / '业绩摘要.docx'
            doc = Document()
            doc.add_paragraph('内地运营商及香港主要竞对关键业绩摘要')
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = 'HKT 2026H1'
            table.cell(0, 1).text = '收入18685百万港元'
            doc.save(report)
            with patch('tts_service._generate_audio_summary_with_llm', return_value='香港电讯业务保持增长。' * 35) as editor:
                build_audio_summary(report)
            self.assertIn('HKT 2026H1 | 收入18685百万港元', editor.call_args.args[0])
            self.assertEqual(editor.call_args.kwargs['report_kind'], 'carrier-performance')

    def test_spoken_script_has_no_written_headings_or_html_entities(self):
        self.assertEqual(normalize_for_speech('结论：香港电讯业务增长。\n摘要：&#x4E2D;国移动推进转型。'),
                         '香港电讯业务增长。 中国移动推进转型')

    def test_decimal_percentage_is_spoken_in_chinese_order(self):
        self.assertEqual(prepare_tts_text("增长8.3%"), "增长百分之八点三")

    def test_multiple_and_negative_percentages(self):
        self.assertEqual(
            prepare_tts_text("增长8.3%，下降17.6%，变动-4.2%"),
            "增长百分之八点三，下降百分之十七点六，变动百分之负四点二",
        )

    def test_integer_percentage(self):
        self.assertEqual(prepare_tts_text("派息率75%"), "派息率百分之七十五")

    def test_3hk_hong_kong_is_not_duplicated(self):
        self.assertEqual(prepare_tts_text("3HK香港业务"), "Three香港业务")
        self.assertEqual(prepare_tts_text("Three香港香港业务"), "Three香港业务")
        self.assertEqual(normalize_for_speech("3HK香港业务"), "Three香港业务")


if __name__ == "__main__":
    unittest.main()
