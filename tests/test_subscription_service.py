import hashlib
import json
import os
import sqlite3
import subprocess
import tempfile
import unittest
from contextlib import closing
from datetime import datetime
from pathlib import Path
from unittest import mock

from docx import Document

from cmhk.services.subscriptions import (
    NEWS_CATEGORY_LABELS,
    NEWS_DIGEST_PREFIX,
    SubscriptionService,
    encode_strategic_news_digest,
    strategic_news_card,
    subscription_confirmation_card,
    subscription_entry_card,
)


class FakeLark:
    def __init__(self, *, job_title="经理", leader_open_id=""):
        self.calls = []
        self.job_title = job_title
        self.leader_open_id = leader_open_id

    def __call__(self, argv, timeout=45):
        self.calls.append(list(argv))
        if "api" in argv and "/open-apis/contact/v3/departments/0/children" in argv:
            payload = {"ok": True, "data": {"has_more": False, "items": [{
                "open_department_id": "od-test123", "name": "战略部", "member_count": 1,
                "leader_user_id": self.leader_open_id,
                "status": {"is_deleted": False},
            }]}}
        elif "api" in argv and "/open-apis/contact/v3/users/find_by_department" in argv:
            payload = {"ok": True, "data": {"has_more": False, "items": [{
                "open_id": "ou_delivery123", "union_id": "on_test123",
                "name": "测试用户", "en_name": "测试用户 Test User", "job_title": self.job_title,
                "enterprise_email": "test.user@hk.chinamobile.com",
                "avatar": {"avatar_72": "https://example.test/avatar.png"},
            }]}}
        elif "api" in argv and "/open-apis/im/v1/chats" in argv:
            payload = {"ok": True, "data": {"has_more": False, "items": [{
                "chat_id": "oc_strategy123", "name": "战略情报群", "description": "每日竞对资讯",
                "chat_mode": "group", "chat_status": "normal", "external": False,
            }, {
                "chat_id": "oc_archived123", "name": "战略旧群", "description": "",
                "chat_mode": "group", "chat_status": "stopped", "external": False,
            }]}}
        elif "+get-user" in argv:
            if "union_id" in argv:
                payload = {"ok": True, "data": {"user": {
                    "open_id": "ou_delivery123", "union_id": "on_test123", "name": "测试用户",
                    "avatar": {"avatar_72": "https://example.test/avatar.png"},
                }}}
            else:
                requested = argv[argv.index("--user-id") + 1]
                payload = {"ok": True, "data": {"user": {
                    "open_id": requested, "union_id": "on_test123", "name": "测试用户",
                    "avatar": {"avatar_72": "https://example.test/avatar.png"},
                }}}
        elif "+messages-send" in argv:
            payload = {"ok": True, "data": {"message_id": "om_test123", "chat_id": "oc_test123"}}
        elif "+messages-mget" in argv:
            payload = {"ok": True, "data": {"items": [{"message_id": "om_test123"}]}}
        else:
            payload = {"ok": True, "data": {}}
        return subprocess.CompletedProcess(argv, 0, json.dumps(payload, ensure_ascii=False), "")


class SubscriptionServiceTests(unittest.TestCase):
    def setUp(self):
        from tests.news_push_fixtures import prepared_assets
        assets = mock.patch('cmhk.services.news_delivery_guard.prepare_news_assets', side_effect=prepared_assets)
        assets.start()
        self.addCleanup(assets.stop)
        from cmhk.services.news_delivery_dedupe import exact_unique
        dedupe = mock.patch("cmhk.services.news_delivery_guard.deduplicate_events",
                            side_effect=lambda items, history, root: (exact_unique(items, history), []))
        dedupe.start()
        self.addCleanup(dedupe.stop)
        editor = mock.patch("cmhk.services.news_digest_editor.prepare_digest",
                            side_effect=lambda items, root: {"items": items, "overview": "测试新闻综述。"})
        editor.start()
        self.addCleanup(editor.stop)
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        (self.root / "config").mkdir()
        (self.root / "config" / "project_monitor.json").write_text(json.dumps({
            "strategic_scan_times": ["03:00", "14:00"],
            "bot": {"profile": "cli_test"},
            "subscriptions": {
                "entry_profile": "cli_test",
                "directory_profile": "org_test",
                "delivery_profile": "org_test",
                "primary_delivery_open_id": "ou_delivery123",
                "confirmation_image_keys": {
                    "cli_test": "img_v3_confirmation_test",
                    "org_test": "img_v3_confirmation_test",
                },
                "preference_updated_image_keys": {
                    "cli_test": "img_v3_preferences_updated_test",
                    "org_test": "img_v3_preferences_updated_test",
                },
                "news_image_keys": {
                    "morning": "img_v3_morning_tea_v2",
                    "afternoon": "img_v3_afternoon_tea_v2",
                },
            },
            "card_actions": {
                "primary_handler_open_id": "ou_test123",
                "primary_handler_expected_name": "测试用户",
            },
            "targets": [
                {"role": "project", "chat_id": "oc_test123", "expected_name": "项目群"},
                {"role": "incident", "chat_id": "oc_incident123", "expected_name": "故障群"},
            ],
        }, ensure_ascii=False), encoding="utf-8")
        self.lark = FakeLark()
        self.service = SubscriptionService(runtime_root=self.root, command_runner=self.lark)

    def tearDown(self):
        self.temp.cleanup()

    def _write_weekly_quality_sidecar(
        self,
        report_path: Path,
        *,
        item_count: int = 4,
        detail_chars: int = 96,
        detail_sentences: int | None = 2,
        report_sha256: str = "",
    ) -> Path:
        payload = {
            "schemaVersion": 1,
            "reportFile": report_path.name,
            "reportSha256": report_sha256 or hashlib.sha256(report_path.read_bytes()).hexdigest(),
            "reportBytes": report_path.stat().st_size,
            "reviewStatus": "passed",
            "generationMode": "normal",
            "limitations": [],
            "qualityWarnings": [],
            "included": item_count,
            "items": [
                {
                    "id": f"W{index:03d}",
                    "detailChars": detail_chars,
                    "detailSentences": detail_sentences,
                }
                for index in range(1, item_count + 1)
            ],
        }
        sidecar_path = Path(str(report_path) + ".quality.json")
        sidecar_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        return sidecar_path

    def _write_weekly_report(self, report_path: Path, *, detail: str = "") -> Path:
        document = Document()
        document.add_paragraph("战略双周报")
        document.add_paragraph(
            detail
            or (
                "测试主体公布经审核的业务进展，披露具体参与范围、执行时间和当前状态。"
                "第二句继续说明可核验的关键数字、后续安排及业务影响，确保正文不是标题重复或网页导航。"
            )
        )
        document.save(report_path)
        return report_path

    def test_card_is_card_2_form_with_three_services(self):
        card = subscription_entry_card(
            image_key="img_v3_subscription_poster",
            recipient_name="Alex LIAO Wang",
            report_schedule={"enabled": True, "days_text": "5、20 日", "time": "08:45"},
            performance_schedule={"enabled": True, "days_text": "10、25 日", "time": "09:30"},
        )
        self.assertEqual(card["schema"], "2.0")
        self.assertTrue(card["config"]["update_multi"])
        self.assertEqual(card["header"]["title"]["content"], "订阅战略情报")
        self.assertNotIn("subtitle", card["header"])
        self.assertNotIn("icon", card["header"])
        self.assertNotIn("text_tag_list", card["header"])
        poster = card["body"]["elements"][0]
        self.assertEqual(poster["tag"], "img")
        self.assertEqual(poster["img_key"], "img_v3_subscription_poster")
        intro = card["body"]["elements"][1]
        self.assertEqual(intro["tag"], "markdown")
        self.assertEqual(
            intro["content"],
            "尊敬的 Alex LIAO Wang，您好！我是战略竞对中心管家小竞。"
            "为帮助战略部宣传和推广战略情报产品，您可以按需选择战略双周报、运营商业绩摘要或战略新闻，"
            "报告按后台设定的月度排期自动生成并推送；战略新闻爬虫每日香港时间 03:00 和 14:00 执行，"
            "个人默认在 08:00 和 18:30 推送，但只有对应爬虫完成审核后才会发送。"
            "感谢您的配合！",
        )

        form = next(item for item in card["body"]["elements"] if item["tag"] == "form")
        self.assertEqual(
            [item["content"] for item in form["elements"] if item["tag"] == "markdown" and item["content"].startswith("**")],
            ["**01 · 选择订阅内容**", "**02 · 报告设置**\n<font color='grey'>适用于战略双周报和运营商业绩摘要。</font>", "**定期推送日期（后台设置）**\n**战略双周报：**每月 5、20 日 08:45（香港时间）\n**运营商业绩摘要：**每月 10、25 日 09:30（香港时间）", "**报告接收方式**", "**03 · 战略新闻设置**\n<font color='grey'>仅订阅战略新闻时生效；以下选项不影响报告推送。</font>", "**感兴趣的战略新闻板块（可多选）**", "**战略新闻频率**\n<font color='grey'>选择每天一次时，只在上午推送，并使用下方第一次时间。</font>", "**新闻地域偏好**", "**每次战略新闻条数**", "**期待收到战略新闻的时间（香港）**\n早间早于08:00、下午早于14:00将自动调整到下限；无效时间使用08:00 / 18:30，成功消息会说明调整结果。"],
        )
        selector = next(item for item in form["elements"] if item["tag"] == "multi_select_static")
        self.assertEqual({item["value"] for item in selector["options"]}, {"weekly", "performance", "news"})
        report_mode = next(item for item in form["elements"] if item.get("name") == "report_mode")
        frequency = next(item for item in form["elements"] if item.get("name") == "news_frequency")
        item_limit = next(item for item in form["elements"] if item.get("name") == "news_item_limit")
        categories = next(item for item in form["elements"] if item.get("name") == "news_categories")
        self.assertEqual({item["value"] for item in report_mode["options"]}, {"pdf", "pdf_audio", "audio"})
        self.assertEqual({item["value"] for item in frequency["options"]}, {"once_daily", "twice_daily"})
        time_pickers = [item for item in form["elements"] if item["tag"] == "picker_time"]
        self.assertEqual([item["initial_time"] for item in time_pickers], ["08:00", "18:30"])
        self.assertTrue(all(not item["required"] for item in time_pickers))
        self.assertEqual({item["value"] for item in item_limit["options"]}, {"5", "10", "15", "20"})
        self.assertEqual(len(categories["options"]), 7)
        self.assertFalse(categories["required"])
        schedule_copy = next(
            item["content"] for item in form["elements"]
            if item.get("tag") == "markdown" and "定期推送日期" in item.get("content", "")
        )
        self.assertIn("战略双周报：**每月 5、20 日 08:45（香港时间）", schedule_copy)
        self.assertIn("运营商业绩摘要：**每月 10、25 日 09:30（香港时间）", schedule_copy)
        self.assertNotIn("周报按后台月度排期自动生成并推送", json.dumps(card, ensure_ascii=False))
        self.assertNotIn("frequency", {item.get("name") for item in form["elements"]})
        button = next(item for item in form["elements"] if item["tag"] == "button")
        self.assertEqual(button["form_action_type"], "submit")
        self.assertEqual(button["text"]["content"], "确认订阅")
        pause = next(item for item in card["body"]["elements"] if item.get("behaviors"))
        self.assertEqual(pause["type"], "text")
        self.assertEqual(pause["behaviors"][0]["value"]["action"], "cmhk_subscription_pause_all_v1")
        footer = card["body"]["elements"][-1]["content"]
        self.assertIn("私聊我", footer)
        self.assertIn("群里 @科创及数智化", footer)
        self.assertIn("加入订阅名单", footer)
        self.assertIn("退出订阅名单", footer)

    def test_server_profiles_can_be_overridden_by_environment(self):
        service = SubscriptionService(
            runtime_root=self.root,
            db_path=self.root / "var" / "subscriptions" / "profile-test.sqlite3",
            environ={
                "CMHK_FEISHU_ENTRY_PROFILE": "server-entry",
                "CMHK_FEISHU_DIRECTORY_PROFILE": "server-directory",
                "CMHK_FEISHU_DELIVERY_PROFILE": "server-delivery",
            },
            command_runner=self.lark,
        )
        self.assertEqual(service.entry_profile, "server-entry")
        self.assertEqual(service.directory_profile, "server-directory")
        self.assertEqual(service.delivery_profile, "server-delivery")

    def test_management_snapshot_exposes_strategic_news_schedule(self):
        self.assertEqual(
            self.service.strategic_news_schedule_snapshot(),
            {
                "service": "news",
                "enabled": True,
                "times": ["03:00", "14:00"],
                "times_text": "03:00 / 14:00",
                "timezone": "Asia/Hong_Kong",
                "timezone_label": "香港时间",
                "dispatch_rule": "个人按本人设定时间推送，且必须等对应爬虫完成",
                "updated_at": self.service.strategic_news_schedule_snapshot()["updated_at"],
            },
        )

    def test_subscription_confirmation_is_compact_card_2_receipt(self):
        card = subscription_confirmation_card(
            image_key="img_v3_confirmation_test",
            display_name="Alex LIAO Wang",
            service_labels="战略双周报、战略新闻",
            report_mode_label="PDF + 单独语音",
            frequency_label="每天两次",
            category_labels="竞对动态、政策监管",
            news_item_limit=15,
        )
        self.assertEqual(card["schema"], "2.0")
        self.assertEqual(card["header"]["template"], "green")
        self.assertEqual(card["header"]["title"]["content"], "订阅已生效")
        self.assertEqual(card["header"]["text_tag_list"][0]["text"]["content"], "已开启")
        self.assertEqual(card["body"]["elements"][0]["tag"], "img")
        self.assertEqual(card["body"]["elements"][0]["img_key"], "img_v3_confirmation_test")
        text = json.dumps(card, ensure_ascii=False)
        self.assertIn("Alex LIAO Wang，设置完成", text)
        self.assertIn("战略双周报、战略新闻", text)
        self.assertIn("每天两次 · 最新 15 条", text)
        self.assertIn("竞对动态、政策监管", text)

    def test_success_callback_sends_interactive_confirmation_card(self):
        self.service.publish_entry_card(target_id="oc_test123", target_type="chat")
        self.lark.calls.clear()
        self.service.handle_card_event({
            "type": "card.action.trigger",
            "action_tag": "button",
            "event_id": "event-card-confirmation",
            "operator_id": "ou_callback123",
            "chat_id": "oc_test123",
            "message_id": "om_test123",
            "form_value": json.dumps({
                "services": ["weekly", "news"],
                "report_mode": "pdf_audio",
                "news_frequency": "twice_daily",
                "news_item_limit": "15",
                "news_categories": ["竞对动态", "政策监管"],
            }),
        })
        send_call = next(call for call in self.lark.calls if "+messages-send" in call)
        self.assertEqual(send_call[send_call.index("--msg-type") + 1], "interactive")
        self.assertNotIn("--markdown", send_call)
        card = json.loads(send_call[send_call.index("--content") + 1])
        self.assertEqual(card["schema"], "2.0")
        self.assertEqual(card["header"]["title"]["content"], "订阅已生效")
        self.assertEqual(card["body"]["elements"][0]["img_key"], "img_v3_confirmation_test")
        self.assertIn("每天两次 · 最新 15 条", json.dumps(card, ensure_ascii=False))
        self.assertNotIn("**", json.dumps(card, ensure_ascii=False))

    def test_preference_update_callback_sends_distinct_feedback_card(self):
        self.service.publish_entry_card(target_id="oc_test123", target_type="chat")
        event = {
            "type": "card.action.trigger",
            "action_tag": "button",
            "operator_id": "ou_callback123",
            "chat_id": "oc_test123",
            "message_id": "om_test123",
        }
        first = self.service.handle_card_event({
            **event,
            "event_id": "event-preference-initial",
            "form_value": json.dumps({
                "services": ["news"],
                "news_categories": ["竞对动态"],
                "news_frequency": "once_daily",
                "news_item_limit": "10",
            }),
        })
        self.assertEqual(first["feedback_kind"], "subscription_started")
        self.lark.calls.clear()

        updated = self.service.handle_card_event({
            **event,
            "event_id": "event-preference-updated",
            "form_value": json.dumps({
                "services": ["news"],
                "news_categories": ["政策监管", "宏观与国际"],
                "news_frequency": "twice_daily",
                "news_item_limit": "15",
                "news_region_preference": "international",
            }),
        })

        self.assertEqual(updated["feedback_kind"], "preference_updated")
        send_call = next(call for call in self.lark.calls if "+messages-send" in call)
        card = json.loads(send_call[send_call.index("--content") + 1])
        self.assertEqual(card["header"]["title"]["content"], "兴趣偏好已更新")
        self.assertEqual(card["header"]["text_tag_list"][0]["text"]["content"], "已更新")
        self.assertEqual(card["body"]["elements"][0]["img_key"], "img_v3_preferences_updated_test")
        self.assertIn("修改成功", json.dumps(card, ensure_ascii=False))
        self.assertEqual(updated["news_region_preference"], "international")
        self.assertEqual(self.service.list_summary()["subscribers"][0]["news_region_preference"], "international")
        self.assertIn("国际新闻优先", json.dumps(card, ensure_ascii=False))

    def test_news_delivery_uses_per_subscriber_schedule_without_global_pause(self):
        self.assertTrue(self.service.strategic_news_schedule_snapshot()["enabled"])
        self.service.update_news_schedule(enabled=False)
        self.assertTrue(self.service.automatic_delivery_enabled("news"))

    def test_form_callback_persists_identity_and_replaces_services(self):
        self.service.publish_entry_card(target_id="oc_test123", target_type="chat")
        first = self.service.handle_card_event({
            "type": "card.action.trigger",
            "action_tag": "button",
            "event_id": "event-1",
            "operator_id": "ou_callback123",
            "chat_id": "oc_test123",
            "message_id": "om_test123",
            "form_value": json.dumps({"services": ["weekly", "news"], "report_mode": "pdf_audio", "news_frequency": "daily", "news_item_limit": "15", "news_categories": ["竞对动态", "政策监管"], "news_delivery_time_morning": "08:15 +0800", "news_delivery_time_afternoon": "18:45 +0800"}),
        })
        self.assertEqual(first["status"], "subscription_saved")
        self.assertEqual(first["services"], ["news", "weekly"])
        self.assertEqual(first["frequency"], "once_daily")
        self.assertEqual(first["news_frequency"], "once_daily")
        self.assertEqual(first["report_cadence"], "biweekly_on_publish")
        self.assertEqual(first["report_mode"], "pdf_audio")
        self.assertEqual(first["news_item_limit"], 15)
        self.assertEqual(first["news_categories"], ["竞对动态", "政策监管"])
        self.assertEqual(first["news_delivery_times"], ["08:15", "18:45"])
        self.service.handle_card_event({
            "type": "card.action.trigger",
            "action_tag": "button",
            "event_id": "event-2",
            "operator_id": "ou_callback123",
            "chat_id": "oc_test123",
            "message_id": "om_test123",
            "form_value": json.dumps({"services": ["performance"], "report_mode": "pdf", "frequency": "weekly"}),
        })
        summary = self.service.list_summary()
        self.assertEqual(summary["active_subscriber_count"], 1)
        self.assertEqual(summary["subscribers"][0]["services"], ["performance"])
        self.assertEqual(summary["subscribers"][0]["frequency"], "once_daily")
        self.assertEqual(summary["subscribers"][0]["report_mode"], "pdf")
        counts = {item["key"]: item["subscriber_count"] for item in summary["services"]}
        self.assertEqual(counts["performance"], 1)

        paused = self.service.handle_card_event({
            "type": "card.action.trigger",
            "action_tag": "button",
            "action_value": json.dumps({"action": "cmhk_subscription_pause_all_v1"}),
            "event_id": "event-3",
            "operator_id": "ou_callback123",
            "chat_id": "oc_test123",
            "message_id": "om_test123",
        })
        self.assertEqual(paused["status"], "subscription_paused")
        self.assertEqual(self.service.list_summary()["active_subscriber_count"], 0)

    def test_real_card_single_select_arrays_are_accepted(self):
        self.service._send_entry_card_to_user("ou_callback123")
        result = self.service.handle_card_event({
            "type": "card.action.trigger",
            "action_tag": "button",
            "event_id": "event-card-array-values",
            "operator_id": "ou_callback123",
            "chat_id": "oc_test123",
            "message_id": "om_test123",
            "form_value": json.dumps({
                "services": ["weekly", "news"],
                "report_mode": ["pdf_audio"],
                "news_frequency": ["twice_daily"],
                "news_item_limit": ["20"],
                "news_categories": ["市场/产品类", "基础设施/网络/技术类"],
            }),
        })
        self.assertEqual(result["status"], "subscription_saved")
        self.assertEqual(result["frequency"], "twice_daily")
        self.assertEqual(result["report_mode"], "pdf_audio")
        self.assertEqual(result["news_item_limit"], 20)
        self.assertEqual(result["news_categories"], ["市场/产品类", "基础设施/网络/技术类"])
        invitation = self.service.list_summary()["invitations"][0]
        self.assertEqual(invitation["message_id"], "om_test123")
        self.assertEqual(invitation["status"], "accepted")
        self.assertTrue(invitation["responded_at"])

    def test_unpublished_subscription_card_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "受控卡片"):
            self.service.handle_card_event({
                "type": "card.action.trigger",
                "action_tag": "button",
                "event_id": "event-forged",
                "operator_id": "ou_callback123",
                "chat_id": "oc_test123",
                "message_id": "om_unpublished123",
                "form_value": json.dumps({"services": ["news"], "frequency": "immediate"}),
            })

    def test_publish_is_whitelisted_and_read_back(self):
        result = self.service.publish_entry_card(target_id="oc_test123", target_type="chat")
        self.assertTrue(result["verified"])
        self.assertTrue(any("+messages-mget" in call for call in self.lark.calls))
        with self.assertRaises(ValueError):
            self.service.publish_entry_card(target_id="oc_incident123", target_type="chat")

    def test_directory_search_returns_avatar_and_adds_controlled_candidate(self):
        refreshed = self.service.refresh_people_directory()
        self.assertEqual(refreshed["people_count"], 1)
        results = self.service.search_people_directory("测试")
        self.assertEqual(results[0]["avatar_url"], "https://example.test/avatar.png")
        self.assertEqual(results[0]["department_names"], ["战略部"])
        self.assertEqual(
            self.service.search_people_directory("test.user@")[0]["display_name"],
            "测试用户",
        )
        self.assertEqual(self.service.invitation_permission_snapshot()["people_count"], 1)
        self.assertEqual(self.service.invitation_permission_snapshot()["status"], "ready")
        added = self.service.add_directory_candidates(["ou_delivery123"])
        self.assertEqual(added["added_count"], 1)
        self.assertEqual(added["candidates"][0]["display_name"], "测试用户")

    def test_group_only_subscriber_inherits_directory_avatar_across_app_ids(self):
        self.service.refresh_people_directory()
        self.service.publish_entry_card(target_id="oc_test123", target_type="chat")
        self.service.handle_card_event({
            "type": "card.action.trigger", "action_tag": "button",
            "event_id": "avatar-group", "operator_id": "ou_callback123",
            "chat_id": "oc_test123", "message_id": "om_test123",
            "form_value": json.dumps({"services": ["news"]}),
        })
        with closing(self.service._connect()) as db, db:
            db.execute("UPDATE subscription_directory_people SET avatar_url='https://s3-imfile.feishucdn.com/avatar.png'")
            db.execute("UPDATE subscription_group_responses SET avatar_url=''")
        person = self.service.list_invite_candidates()[0]
        self.assertEqual(person["directory_open_id"], "ou_delivery123")
        self.assertEqual(person["department_names"], ["战略部"])
        self.assertEqual(person["avatar_url"], "https://s3-imfile.feishucdn.com/avatar.png")
        self.assertEqual(self.service.avatar_source_url("ou_callback123"), person["avatar_url"])

    def test_directory_refresh_does_not_infer_position_from_org_relationship(self):
        self.service.command_runner = FakeLark(job_title="", leader_open_id="ou_delivery123")
        self.service.refresh_people_directory()
        self.assertEqual(self.service.search_people_directory("测试")[0]["job_title"], "")

        self.service.command_runner = FakeLark(job_title="经理", leader_open_id="ou_someone_else")
        self.service.refresh_people_directory()
        self.assertEqual(self.service.search_people_directory("测试")[0]["job_title"], "经理")

    def test_chat_search_returns_only_visible_normal_group_matches(self):
        results = self.service.search_chat_directory("战略")
        self.assertEqual([item["name"] for item in results], ["战略情报群"])
        self.assertEqual(results[0]["chat_id"], "oc_strategy123")
        self.assertEqual(self.service.search_chat_directory("不存在"), [])

    def test_admin_invite_target_auto_detects_group_and_sends_group_compatible_card(self):
        with self.assertRaisesRegex(ValueError, "二次确认"):
            self.service.invite_target("oc_strategy123")
        sent = self.service.invite_target("oc_strategy123", confirm_invite=True)
        self.assertEqual(sent["target_type"], "chat")
        self.assertEqual(sent["target_id"], "oc_strategy123")
        self.assertEqual(sent["target_name"], "战略情报群")
        send_call = next(call for call in self.lark.calls if "+messages-send" in call)
        self.assertEqual(send_call[send_call.index("--chat-id") + 1], "oc_strategy123")
        card = json.loads(send_call[send_call.index("--content") + 1])
        self.assertTrue(card["config"]["update_multi"])
        self.assertNotIn("**", json.dumps(card, ensure_ascii=False))
        group_invitation = self.service.list_summary()["group_invitations"][0]
        self.assertEqual(group_invitation["target_name"], "战略情报群")
        self.assertEqual(group_invitation["status"], "verified")
        self.assertEqual(group_invitation["response_count"], 0)

    def test_two_people_can_save_different_preferences_from_one_group_card(self):
        self.service.publish_entry_card(target_id="oc_test123", target_type="chat")

        def identity(open_id, *, source_profile=""):
            return {
                "display_name": open_id,
                "callback_open_id": open_id,
                "union_id": f"on_{open_id[3:]}",
                "open_id": open_id,
                "source_profile": source_profile,
                "avatar_url": "",
                "job_title": "",
            }
        with mock.patch.object(self.service, "resolve_user", side_effect=identity):
            for index, (operator_id, categories, item_limit) in enumerate((
                ("ou_persona123", ["竞对动态"], "5"),
                ("ou_personb123", ["政策监管"], "20"),
            ), start=1):
                self.service.handle_card_event({
                    "type": "card.action.trigger",
                    "action_tag": "button",
                    "event_id": f"event-group-{index}",
                    "operator_id": operator_id,
                    "chat_id": "oc_test123",
                    "message_id": "om_test123",
                    "form_value": json.dumps({
                        "services": ["news"],
                        "news_frequency": "once_daily",
                        "report_mode": "pdf",
                        "news_item_limit": item_limit,
                        "news_categories": categories,
                    }),
                })
        subscribers = {item["open_id"]: item for item in self.service.list_summary()["subscribers"]}
        self.assertEqual(set(subscribers), {"ou_persona123", "ou_personb123"})
        self.assertEqual(subscribers["ou_persona123"]["news_categories"], ["竞对动态"])
        self.assertEqual(subscribers["ou_persona123"]["news_item_limit"], 5)
        self.assertEqual(subscribers["ou_personb123"]["news_categories"], ["政策监管"])
        self.assertEqual(subscribers["ou_personb123"]["news_item_limit"], 20)
        self.assertEqual(subscribers["ou_persona123"]["preference_source"], "group_card")
        self.assertEqual(subscribers["ou_personb123"]["preference_message_id"], "om_test123")
        self.service.update_subscriber("ou_persona123", services=["weekly"], news_item_limit=20)
        restored = self.service.reset_subscriber("ou_persona123")
        self.assertEqual(restored["services"], ["news"])
        self.assertEqual(restored["news_categories"], ["竞对动态"])
        self.assertEqual(restored["news_item_limit"], 5)
        group_invitation = self.service.list_summary()["group_invitations"][0]
        self.assertEqual(group_invitation["status"], "responded")
        self.assertEqual(group_invitation["response_count"], 2)
        self.assertEqual(
            {item["display_name"] for item in group_invitation["responses"]},
            {"ou_persona123", "ou_personb123"},
        )
        with closing(self.service._connect()) as db:
            db.execute("DELETE FROM subscription_group_responses")
            db.commit()
        reloaded = SubscriptionService(runtime_root=self.root, command_runner=self.lark)
        self.assertEqual(reloaded.list_summary()["group_invitations"][0]["response_count"], 2)

    def test_user_preference_changes_are_kept_in_submission_history(self):
        self.service.publish_entry_card(target_id="oc_test123", target_type="chat")

        def identity(open_id, *, source_profile=""):
            return {
                "display_name": "测试用户",
                "callback_open_id": open_id,
                "union_id": "on_persona123",
                "open_id": open_id,
                "source_profile": source_profile,
                "avatar_url": "",
                "job_title": "",
            }

        forms = (
            {
                "services": ["news"],
                "news_frequency": "once_daily",
                "report_mode": "pdf",
                "news_item_limit": "5",
                "news_categories": ["竞对动态"],
                "news_delivery_time_morning": "08:00",
                "news_delivery_time_afternoon": "18:30",
            },
            {
                "services": ["weekly", "news"],
                "news_frequency": "twice_daily",
                "report_mode": "pdf_audio",
                "news_item_limit": "20",
                "news_categories": ["政策监管"],
                "news_delivery_time_morning": "09:00",
                "news_delivery_time_afternoon": "19:00",
            },
        )
        with mock.patch.object(self.service, "resolve_user", side_effect=identity):
            for index, form in enumerate(forms, start=1):
                self.service.handle_card_event({
                    "type": "card.action.trigger",
                    "action_tag": "button",
                    "event_id": f"event-preference-{index}",
                    "operator_id": "ou_persona123",
                    "chat_id": "oc_test123",
                    "message_id": "om_test123",
                    "form_value": json.dumps(form),
                })

        summary = self.service.list_summary()
        response = summary["group_invitations"][0]["responses"][0]
        submissions = response["submissions"]
        self.assertEqual(len(submissions), 2)
        self.assertTrue(submissions[1]["is_initial"])
        self.assertEqual(submissions[1]["changes"], [])
        self.assertFalse(submissions[0]["is_initial"])
        changes = {item["field"]: item for item in submissions[0]["changes"]}
        self.assertEqual(changes["frequency"]["before"], "每天一次（上午）")
        self.assertEqual(changes["frequency"]["after"], "每天两次")
        self.assertEqual(changes["news_item_limit"]["before"], "5 条")
        self.assertEqual(changes["news_item_limit"]["after"], "20 条")
        self.assertEqual(changes["news_categories"]["before"], "竞对动态")
        self.assertEqual(changes["news_categories"]["after"], "政策监管")
        self.assertEqual(len(summary["preference_submissions"]), 2)

        self.service.update_subscriber(
            "ou_persona123", services=["weekly"], report_mode="audio",
            news_item_limit=10, news_categories=["公司动态"],
        )
        self.assertEqual(len(self.service.list_summary()["preference_submissions"]), 2)

    def test_repeated_group_cards_accumulate_unique_people_and_complete_profiles(self):
        self.service.refresh_people_directory()
        self.service.publish_entry_card(target_id="oc_test123", target_type="chat")
        self.service.handle_card_event({
            "type": "card.action.trigger",
            "action_tag": "button",
            "event_id": "event-group-first",
            "operator_id": "ou_callback123",
            "chat_id": "oc_test123",
            "message_id": "om_test123",
            "form_value": json.dumps({"services": ["news"]}),
        })
        with closing(self.service._connect()) as db, db:
            db.execute(
                """INSERT INTO subscription_entry_cards(
                       message_id, target_type, target_id, target_name, chat_id, source_profile, created_at
                   ) VALUES('om_second123', 'chat', 'oc_test123', '项目群', 'oc_test123', 'cli_test',
                            '2026-09-09T18:00:00+08:00')"""
            )
            db.execute(
                """INSERT INTO subscription_directory_people(
                       directory_open_id, union_id, display_name, en_name, avatar_url, job_title,
                       department_names, source_profile, active, synced_at
                   ) VALUES('ou_directory456', 'on_other456', '朱子旭', 'Red ZHU Zixu',
                            'https://example.test/other.png', '高级经理', '["战略部"]',
                            'org_test', 1, '2026-09-09T18:01:00+08:00')"""
            )
            for values in (
                (
                    "ou_callback123", "ou_delivery123", "on_test123", "测试用户",
                    "2026-09-09T18:02:00+08:00",
                ),
                (
                    "ou_callback456", "ou_delivery456", "on_other456", "朱子旭",
                    "2026-09-09T18:03:00+08:00",
                ),
            ):
                db.execute(
                    """INSERT INTO subscription_group_responses(
                           message_id, chat_id, callback_open_id, delivery_open_id, union_id,
                           display_name, avatar_url, source_profile, department_names, job_title,
                           status, responded_at, updated_at
                       ) VALUES('om_second123', 'oc_test123', ?, ?, ?, ?, '', 'cli_test',
                                '[]', '', 'accepted', ?, ?)""",
                    (*values, values[-1]),
                )
            db.execute(
                """INSERT INTO subscription_invitations(
                       callback_open_id, delivery_open_id, union_id, display_name,
                       source_profile, avatar_url, message_id, chat_id, status,
                       invited_by, sent_at, updated_at
                   ) VALUES('ou_callback123', 'ou_delivery123', 'on_test123', '测试用户',
                            'cli_test', '', 'om_personal_pending123', 'oc_test123', 'pending',
                            'local_admin', '2026-09-09T18:04:00+08:00',
                            '2026-09-09T18:04:00+08:00')"""
            )

        groups = self.service.list_summary()["group_invitations"]
        self.assertEqual(len(groups), 1)
        group = groups[0]
        self.assertEqual(group["message_count"], 2)
        self.assertEqual(group["response_count"], 2)
        self.assertEqual(group["accepted_count"], 2)
        people = {item["union_id"]: item for item in group["responses"]}
        self.assertEqual(people["on_test123"]["department_names"], ["战略部"])
        self.assertEqual(people["on_test123"]["job_title"], "经理")
        self.assertEqual(people["on_other456"]["en_name"], "Red ZHU Zixu")
        self.assertEqual(people["on_other456"]["department_names"], ["战略部"])
        self.assertEqual(
            sum(item["callback_open_id"] == "ou_callback123" for item in group["responses"]),
            1,
        )
        candidate = next(
            item for item in self.service.list_invite_candidates()
            if item["union_id"] == "on_test123"
        )
        self.assertEqual(candidate["latest_invitation"]["status"], "accepted")
        self.assertEqual(candidate["latest_invitation"]["response_source"], "group")

    def test_admin_edit_survives_reload_and_reset_restores_only_target(self):
        self.service.save_subscriptions("ou_persona123", "甲", ["news"], news_categories=["竞对动态"], news_item_limit=5)
        self.service.save_subscriptions("ou_personb123", "乙", ["weekly"])
        self.service.update_subscriber("ou_persona123", services=["weekly", "news"], news_categories=["政策监管"], news_item_limit=20, report_mode="audio", frequency="twice_daily", status="paused")
        reloaded = SubscriptionService(runtime_root=self.root, command_runner=self.lark)
        rows = {r["open_id"]: r for r in reloaded.list_summary()["subscribers"]}
        self.assertEqual(rows["ou_persona123"]["news_item_limit"], 20)
        self.assertEqual(rows["ou_persona123"]["status"], "paused")
        self.assertEqual(rows["ou_persona123"]["default_preferences"]["services"], ["news"])
        self.assertEqual(rows["ou_persona123"]["default_preferences"]["news_categories"], ["竞对动态"])
        self.assertEqual(rows["ou_persona123"]["default_preferences"]["status"], "active")
        restored = reloaded.reset_subscriber("ou_persona123")
        self.assertEqual(restored["services"], ["news"])
        self.assertEqual(restored["news_categories"], ["竞对动态"])
        self.assertEqual(restored["news_item_limit"], 5)
        self.assertEqual(restored["status"], "active")
        self.assertEqual(restored["report_mode"], "pdf")
        self.assertEqual(restored["frequency"], "once_daily")
        rows = {r["open_id"]: r for r in reloaded.list_summary()["subscribers"]}
        self.assertEqual(rows["ou_personb123"]["services"], ["weekly"])

    def test_personal_card_still_rejects_a_different_operator(self):
        self.service._send_entry_card_to_user("ou_invited123")
        with self.assertRaisesRegex(ValueError, "受邀人不一致"):
            self.service.handle_card_event({
                "type": "card.action.trigger",
                "action_tag": "button",
                "event_id": "event-wrong-person",
                "operator_id": "ou_someoneelse123",
                "chat_id": "oc_test123",
                "message_id": "om_test123",
                "form_value": json.dumps({
                    "services": ["news"],
                    "news_frequency": "once_daily",
                    "report_mode": "pdf",
                    "news_item_limit": "10",
                    "news_categories": ["竞对动态"],
                }),
            })

    def test_invite_is_selected_only_and_callback_updates_result(self):
        self.service.refresh_people_directory()
        self.service.add_directory_candidates(["ou_delivery123"])
        with self.assertRaisesRegex(ValueError, "二次确认"):
            self.service.invite_users(["ou_delivery123"])
        sent = self.service.invite_users(["ou_delivery123"], confirm_invite=True)
        self.assertEqual(sent["sent_count"], 1)
        send_call = next(call for call in self.lark.calls if "+messages-send" in call)
        card = json.loads(send_call[send_call.index("--content") + 1])
        intro = next(item for item in card["body"]["elements"] if item.get("tag") == "markdown")
        self.assertTrue(intro["content"].startswith("尊敬的 Test User，您好！"))
        self.assertTrue(intro["content"].endswith("感谢您的配合！"))
        self.assertEqual(self.service.list_summary()["invitations"][0]["status"], "pending")
        accepted = self.service.handle_card_event({
            "type": "card.action.trigger",
            "action_tag": "button",
            "event_id": "event-invite-accept",
            "operator_id": "ou_delivery123",
            "chat_id": "oc_test123",
            "message_id": "om_test123",
            "form_value": json.dumps({"services": ["news"], "frequency": "daily"}),
        })
        self.assertEqual(accepted["status"], "subscription_saved")
        invitation = self.service.list_summary()["invitations"][0]
        self.assertEqual(invitation["status"], "accepted")
        self.assertTrue(invitation["responded_at"])

    def test_invite_rejects_person_outside_controlled_candidates(self):
        with self.assertRaisesRegex(ValueError, "受控名单"):
            self.service.invite_users(["ou_unknown123"], confirm_invite=True)

    def test_news_test_push_is_logged_and_verified(self):
        self.service.save_subscriptions("ou_delivery123", "测试用户", ["news"])
        result = self.service.push(
            service="news",
            mode="text",
            title="真实新闻",
            body="经审核的新闻正文",
            test_open_id="ou_test123",
        )
        self.assertEqual(result["verified_count"], 1)
        delivery = self.service.list_summary()["deliveries"][0]
        self.assertEqual(delivery["status"], "verified")
        self.assertEqual(delivery["service"], "news")
        self.assertEqual(delivery["recipient_name"], "测试用户")
        self.assertEqual(delivery["recipient_open_id"], "ou_delivery123")
        send_call = next(call for call in self.lark.calls if "+messages-send" in call)
        self.assertEqual(send_call[send_call.index("--msg-type") + 1], "interactive")
        self.assertEqual(send_call[send_call.index("--profile") + 1], "org_test")
        self.assertEqual(send_call[send_call.index("--user-id") + 1], "ou_delivery123")
        card = json.loads(send_call[send_call.index("--content") + 1])
        self.assertEqual(card["schema"], "2.0")
        self.assertEqual(card["header"]["title"]["content"], "真实新闻")
        self.assertIn("经审核的新闻正文", json.dumps(card, ensure_ascii=False))

    def test_delivery_summary_returns_complete_history_by_default(self):
        with closing(self.service._connect()) as db, db:
            db.executemany(
                """INSERT INTO deliveries(batch_id, open_id, service, mode, content_ref, status, message_ids, error, created_at)
                   VALUES(?, ?, 'news', 'text', ?, 'verified', '[]', '', ?)""",
                [
                    (f"batch-{index}", "ou_delivery123", f"history-{index}", f"2026-08-{(index % 28) + 1:02d}T08:00:00+08:00")
                    for index in range(85)
                ],
            )

        summary = self.service.list_summary()

        self.assertEqual(len(summary["deliveries"]), 85)
        self.assertEqual(summary["delivery_history"]["total"], 85)
        self.assertEqual(summary["delivery_history"]["returned"], 85)

    def test_strategic_news_card_uses_clean_personal_subscription_format(self):
        items = [{
            "title": f"新闻 {index}", "summary": f"摘要 {index}", "category": "竞对动态",
            "region": "香港本地", "source": "测试来源", "published_at": "2026-08-19T09:00:00+08:00",
            "source_url": f"https://example.test/{index}",
            "image_key": f"img_article_{index}", "image_kind": "source",
            "image_source_url": f"https://example.test/{index}.jpg",
        } for index in range(1, 7)]
        card = strategic_news_card(
            title="  CMHK战略订阅｜6条新闻  ",
            body=encode_strategic_news_digest(items),
            published_at="2026-08-19T10:00:00+08:00",
            image_key="img_v3_morning_tea_v2",
        )
        self.assertNotIn("subtitle", card["header"])
        self.assertEqual(card["schema"], "2.0")
        self.assertEqual(card["body"]["elements"][0]["img_key"], "img_v3_morning_tea_v2")
        text = json.dumps(card, ensure_ascii=False)
        self.assertNotIn("今日核心看点", text)
        self.assertNotIn("今日关键信号", text)
        self.assertNotIn("重点涉及", text)
        self.assertNotIn("01｜", text)
        self.assertNotIn("06｜", text)
        self.assertIn("新闻 6", text)
        self.assertNotIn("###", text)
        self.assertNotIn('AI解读', text)
        self.assertNotIn('feishu.cn/docx', text)
        group = next(e['columns'][0] for e in card['body']['elements'] if e['tag'] == 'column_set')
        rows = [e for e in group['elements'] if e['tag'] == 'interactive_container']
        self.assertEqual(len(rows), 6)
        for index, row in enumerate(rows, 1):
            self.assertEqual(row['behaviors'][0]['default_url'], f'https://example.test/{index}')
            columns = row['elements'][0]['columns']
            self.assertEqual(columns[0]['elements'][1]['content'], f'摘要 {index}')
            self.assertFalse(columns[1]['elements'][0]['preview'])
        self.assertEqual(sum(e['tag'] == 'hr' for e in group['elements']), 5)

    def test_strategic_news_card_groups_categories_with_distinct_backgrounds(self):
        card = strategic_news_card(
            title="CMHK战略下午茶订阅｜2026年08月22日",
            body=encode_strategic_news_digest([
                {"source_url":"https://publisher.example/1", "image_key":"img_test", "title": "竞对新闻", "category": "竞对动态", "region": "香港本地"},
                {"source_url":"https://publisher.example/2", "image_key":"img_test", "title": "政策新闻", "category": "政策监管", "region": "香港本地"},
                {"source_url":"https://publisher.example/3", "image_key":"img_test", "title": "产品新闻", "category": "市场/产品类", "region": "国际/行业"},
            ]),
        )
        groups = [e["columns"][0] for e in card["body"]["elements"] if e["tag"] == "column_set" and "background_style" in e["columns"][0]]
        self.assertEqual(len(groups), 3)
        self.assertEqual(len({g["background_style"] for g in groups}), 3)
        self.assertIn("市场与产品", json.dumps(groups, ensure_ascii=False))
        self.assertNotIn("国际/行业", json.dumps(groups, ensure_ascii=False))
        self.assertEqual(sum(e["tag"] == "hr" for e in card["body"]["elements"]), 1)

    def test_report_test_push_sends_pdf_and_reads_it_back(self):
        from cmhk.reporting.pdf_preview import pdf_preview_path

        report = self.root / "测试周报.docx"
        self._write_weekly_report(report)
        self._write_weekly_quality_sidecar(report)
        pdf = pdf_preview_path(report, self.root / "web" / "static" / "report-previews")
        pdf.parent.mkdir(parents=True)
        pdf.write_bytes(b"%PDF-1.7\n")
        result = self.service.push(
            service="weekly",
            mode="pdf",
            path=report.name,
            test_open_id="ou_test123",
        )
        self.assertEqual(result["verified_count"], 1)
        send_call = next(call for call in self.lark.calls if "--file" in call)
        self.assertEqual(
            send_call[send_call.index("--file") + 1],
            "var/subscriptions/outbound/CMHK_战略双周报_测试周报.pdf",
        )
        named_pdf = self.root / "var" / "subscriptions" / "outbound" / "CMHK_战略双周报_测试周报.pdf"
        self.assertEqual(named_pdf.read_bytes(), pdf.read_bytes())

    def test_explicit_user_edited_weekly_copy_uses_audited_formal_source(self):
        from cmhk.reporting.pdf_preview import pdf_preview_path

        source = self.root / "正式周报.docx"
        edited = self.root / "正式周报（编辑稿）.docx"
        self._write_weekly_report(source)
        self._write_weekly_quality_sidecar(source)
        self._write_weekly_report(edited, detail="这是经过页面编辑器人工修订的完整周报正文，仍然保留正式源文件的审计关系。")
        metadata_path = self.root / "data" / "reporting" / "report_file_metadata.json"
        metadata_path.parent.mkdir(parents=True)
        metadata_path.write_text(json.dumps({
            edited.name: {
                "isEdited": True,
                "reportType": "weekly",
                "editorRevision": 1,
                "sourcePath": source.name,
                "sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            }
        }, ensure_ascii=False), encoding="utf-8")
        preview = pdf_preview_path(edited, self.root / "web" / "static" / "report-previews")
        preview.parent.mkdir(parents=True)
        preview.write_bytes(b"%PDF-1.7\n")

        with self.assertRaisesRegex(RuntimeError, "质量审计"):
            self.service.push(
                service="weekly", mode="pdf", path=edited.name, test_open_id="ou_test123"
            )
        result = self.service.push(
            service="weekly",
            mode="pdf",
            path=edited.name,
            test_open_id="ou_test123",
            allow_user_edited=True,
        )

        self.assertEqual(result["verified_count"], 1)
        self.assertTrue(any("--file" in call for call in self.lark.calls))

    def test_manual_selection_delivers_legacy_and_external_edits_without_generated_audit(self):
        from cmhk.reporting.pdf_preview import pdf_preview_path

        for name in ("历史周报 (2).docx", "8月30日周报 (3)（编辑稿）.docx"):
            with self.subTest(name=name):
                report = self.root / name
                self._write_weekly_report(report)
                preview = pdf_preview_path(report, self.root / "web" / "static" / "report-previews")
                preview.parent.mkdir(parents=True, exist_ok=True)
                preview.write_bytes(b"%PDF-1.7\n")
                with self.assertRaisesRegex(RuntimeError, "质量审计"):
                    self.service.push(service="weekly", mode="pdf", path=name, test_open_id="ou_test123")
                result = self.service.push(
                    service="weekly", mode="pdf", path=name, test_open_id="ou_test123",
                    manual_report_selection=True,
                )
                self.assertEqual(result["verified_count"], 1)
                self.assertFalse(Path(str(report) + ".quality.json").exists())

    def test_manual_selection_rejects_broken_empty_and_escaping_documents(self):
        broken = self.root / "损坏.docx"
        broken.write_bytes(b"not a word document")
        empty = self.root / "空白.docx"
        Document().save(empty)
        for path, reason in ((broken.name, "无法读取"), (empty.name, "正文为空"), ("../outside.docx", "路径无效")):
            with self.subTest(path=path), self.assertRaisesRegex(ValueError, reason):
                self.service.push(
                    service="weekly", mode="pdf", path=path, test_open_id="ou_test123",
                    manual_report_selection=True,
                )
        self.assertFalse(any("+messages-send" in call for call in self.lark.calls))

    def test_user_edited_weekly_flag_cannot_bypass_editor_provenance(self):
        report = self.root / "伪编辑稿.docx"
        self._write_weekly_report(report)

        with self.assertRaisesRegex(RuntimeError, "编辑记录"):
            self.service.push(
                service="weekly",
                mode="pdf",
                path=report.name,
                test_open_id="ou_test123",
                allow_user_edited=True,
            )

    def test_reports_reject_non_pdf_delivery_modes(self):
        report = self.root / "测试周报.docx"
        self._write_weekly_report(report)
        self._write_weekly_quality_sidecar(report)
        with self.assertRaisesRegex(ValueError, "只支持 PDF"):
            self.service.push(service="weekly", mode="text", path=report.name, test_open_id="ou_test123")

    def test_report_audio_is_a_separate_message_after_pdf(self):
        from cmhk.reporting.pdf_preview import pdf_preview_path

        report = self.root / "语音周报.docx"
        self._write_weekly_report(report)
        self._write_weekly_quality_sidecar(report)
        pdf = pdf_preview_path(report, self.root / "web" / "static" / "report-previews")
        pdf.parent.mkdir(parents=True)
        pdf.write_bytes(b"%PDF-1.7\n")
        audio = self.root / "audio" / "语音周报.opus"
        audio.parent.mkdir()
        audio.write_bytes(b"OggS")
        result = self.service.push(
            service="weekly",
            mode="pdf_audio",
            path=report.name,
            test_open_id="ou_test123",
        )
        self.assertEqual(len(result["results"][0]["message_ids"]), 2)
        sends = [call for call in self.lark.calls if "+messages-send" in call]
        self.assertIn("--file", sends[0])
        self.assertIn("--audio", sends[1])
        self.assertEqual(
            sends[0][sends[0].index("--file") + 1],
            "var/subscriptions/outbound/CMHK_战略双周报_语音周报.pdf",
        )
        self.assertEqual(
            sends[1][sends[1].index("--audio") + 1],
            "var/subscriptions/outbound/CMHK_战略双周报_语音周报_音频.opus",
        )

    def test_report_audio_accepts_tts_safe_name_for_numbered_report(self):
        from cmhk.reporting.pdf_preview import pdf_preview_path

        report = self.root / "语音周报 (3).docx"
        self._write_weekly_report(report)
        self._write_weekly_quality_sidecar(report)
        pdf = pdf_preview_path(report, self.root / "web" / "static" / "report-previews")
        pdf.parent.mkdir(parents=True)
        pdf.write_bytes(b"%PDF-1.7\n")
        audio = self.root / "audio" / "语音周报 3.opus"
        audio.parent.mkdir()
        audio.write_bytes(b"OggS-safe-name")

        result = self.service.push(
            service="weekly",
            mode="pdf_audio",
            path=report.name,
            test_open_id="ou_test123",
        )

        self.assertEqual(len(result["results"][0]["message_ids"]), 2)
        audio_send = next(
            call for call in self.lark.calls
            if "+messages-send" in call and "--audio" in call
        )
        sent_audio = self.root / audio_send[audio_send.index("--audio") + 1]
        self.assertEqual(sent_audio.read_bytes(), b"OggS-safe-name")

    def test_missing_audio_fails_before_sending_pdf(self):
        from cmhk.reporting.pdf_preview import pdf_preview_path

        report = self.root / "缺语音周报.docx"
        self._write_weekly_report(report)
        self._write_weekly_quality_sidecar(report)
        pdf = pdf_preview_path(report, self.root / "web" / "static" / "report-previews")
        pdf.parent.mkdir(parents=True)
        pdf.write_bytes(b"%PDF-1.7\n")

        result = self.service.push(
            service="weekly",
            mode="pdf_audio",
            path=report.name,
            test_open_id="ou_test123",
        )

        self.assertEqual(result["results"][0]["status"], "failed")
        self.assertIn("尚无可推送语音", result["results"][0]["error"])
        self.assertFalse(any("+messages-send" in call for call in self.lark.calls))

    def test_bulk_report_audio_respects_each_subscriber_preference(self):
        from cmhk.reporting.pdf_preview import pdf_preview_path

        report = self.root / "偏好周报.docx"
        self._write_weekly_report(report)
        self._write_weekly_quality_sidecar(report)
        pdf = pdf_preview_path(report, self.root / "web" / "static" / "report-previews")
        pdf.parent.mkdir(parents=True)
        pdf.write_bytes(b"%PDF-1.7\n")
        audio = self.root / "audio" / "偏好周报.opus"
        audio.parent.mkdir()
        audio.write_bytes(b"OggS")
        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["weekly"], report_mode="pdf", frequency="daily"
        )
        pdf_only = self.service.push(
            service="weekly", mode="pdf_audio", path=report.name, confirm_bulk=True
        )
        self.assertEqual(pdf_only["results"][0]["mode"], "pdf")
        self.assertEqual(pdf_only["results"][0]["frequency"], "immediate")
        self.assertEqual(pdf_only["verified_count"], 1)
        self.assertEqual(pdf_only["queued_count"], 0)
        self.assertEqual(len(pdf_only["results"][0]["message_ids"]), 1)
        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["weekly"], report_mode="pdf_audio"
        )
        with_audio = self.service.push(
            service="weekly", mode="pdf_audio", path=report.name, confirm_bulk=True
        )
        self.assertEqual(with_audio["results"][0]["mode"], "pdf_audio")
        self.assertEqual(len(with_audio["results"][0]["message_ids"]), 2)
        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["weekly"], report_mode="audio"
        )
        audio_only = self.service.push(
            service="weekly", mode="pdf_audio", path=report.name, confirm_bulk=True
        )
        self.assertEqual(audio_only["results"][0]["mode"], "audio")
        self.assertEqual(len(audio_only["results"][0]["message_ids"]), 1)

    def test_bulk_push_requires_explicit_confirmation(self):
        self.service.save_subscriptions("ou_delivery123", "测试用户", ["news"])
        with self.assertRaisesRegex(ValueError, "二次确认"):
            self.service.push(service="news", mode="text", title="新闻", body="正文")

    def test_manual_push_can_target_one_active_subscriber_without_bulk_confirmation(self):
        self.service.save_subscriptions("ou_delivery123", "测试用户", ["news"])
        result = self.service.push(
            service="news",
            mode="text",
            title="手动推送",
            body="只发送给指定订阅者",
            target_open_id="ou_delivery123",
        )
        self.assertEqual(result["recipient_count"], 1)
        self.assertEqual(result["results"][0]["open_id"], "ou_delivery123")
        self.assertEqual(result["verified_count"], 1)
        with self.assertRaisesRegex(ValueError, "未启用此项服务"):
            self.service.push(
                service="weekly",
                mode="pdf",
                path="missing.docx",
                target_open_id="ou_delivery123",
            )

    @mock.patch.object(SubscriptionService, "_deliver_one", return_value=["om_clock_test"])
    def test_news_dispatches_after_crawler_completion_with_daily_limits(self, delivery):
        self.service.save_subscriptions(
            "ou_delivery123",
            "测试用户",
            ["news"],
            frequency="once_daily",
        )
        self.service.update_news_schedule(enabled=True)
        item = {"title": "爬虫新闻", "summary": "已完成审核", "source_url": "https://example.test/news"}
        morning = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@07:00",
            slot_label="晨间扫描",
            items=[item],
        )
        afternoon = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@15:00",
            slot_label="午后扫描",
            items=[item],
        )
        self.assertEqual(morning["queued_count"], 1)
        self.assertEqual(afternoon["skipped_count"], 1)
        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["news"], frequency="twice_daily"
        )
        next_morning = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-02@07:00", slot_label="晨间扫描", items=[item]
        )
        next_afternoon = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-02@15:00", slot_label="午后扫描", items=[item]
        )
        self.assertEqual(next_morning["queued_count"], 1)
        self.assertEqual(next_afternoon["queued_count"], 1)
        flushed = self.service.flush_due(prepared_news_only=False, now=datetime.fromisoformat("2099-01-02T19:00:00+08:00"))
        self.assertEqual(flushed["verified_count"], 3)
        self.assertEqual(delivery.call_count, 3)

    @mock.patch.object(SubscriptionService, "_deliver_one", return_value=["om_clock_test"])
    def test_personal_news_waits_for_late_crawl_completion(self, delivery):
        self.service.save_subscriptions(
            "ou_delivery123",
            "测试用户",
            ["news"],
            frequency="twice_daily",
            news_delivery_times=["08:00", "18:30"],
        )
        queued = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@03:00",
            slot_label="晨间扫描",
            items=[{"title": "迟完成新闻"}],
            completed_at="2099-01-01T09:17:00+08:00",
        )

        self.assertEqual(queued["queued_count"], 1)
        self.assertEqual(queued["results"][0]["due_at"], "2099-01-01T09:17:00+08:00")
        before = self.service.flush_due(prepared_news_only=False,
            now=datetime.fromisoformat("2099-01-01T09:16:59+08:00")
        )
        self.assertEqual(before["processed_count"], 0)
        self.assertFalse(any("+messages-send" in call for call in self.lark.calls))
        at_completion = self.service.flush_due(prepared_news_only=False,
            now=datetime.fromisoformat("2099-01-01T09:17:00+08:00")
        )
        self.assertEqual(at_completion["verified_count"], 1)

    def test_personal_news_keeps_original_crawl_day_when_completion_crosses_midnight(self):
        self.service.save_subscriptions(
            "ou_delivery123",
            "测试用户",
            ["news"],
            frequency="once_daily",
            news_delivery_times=["08:00", "18:30"],
        )
        delayed = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@03:00",
            slot_label="晨间扫描",
            items=[{
                "title": "1月1日迟到新闻",
                "summary": "跨日完成但仍属于原爬虫日。",
                "published_at": "2099-01-01T12:00:00+08:00",
                "category": "公司动态",
            }],
            completed_at="2099-01-02T00:07:00+08:00",
        )

        self.assertEqual(delayed["queued_count"], 1)
        self.assertEqual(delayed["results"][0]["due_at"], "2099-01-02T00:07:00+08:00")
        with mock.patch("cmhk.services.news_delivery_guard.datetime") as clock:
            clock.now.return_value = datetime.fromisoformat("2099-01-02T00:07:00+08:00")
            sent = self.service.flush_due(prepared_news_only=False, now=clock.now.return_value)
        self.assertEqual(sent["verified_count"], 1)
        send = next(call for call in self.lark.calls if "+messages-send" in call)
        card = json.loads(send[send.index("--content") + 1])
        self.assertEqual(card["header"]["title"]["content"], "CMHK战略早茶订阅｜2099年01月01日")
        sent_text = json.dumps(card, ensure_ascii=False)
        self.assertIn("interactive_container", sent_text)
        self.assertNotIn("AI解读", sent_text)

        next_day = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-02@03:00",
            slot_label="晨间扫描",
            items=[{"title": "1月2日新闻"}],
            completed_at="2099-01-02T07:00:00+08:00",
        )
        self.assertEqual(next_day["queued_count"], 1)
        self.assertEqual(next_day["results"][0]["due_at"], "2099-01-02T08:00:00+08:00")

    def test_news_claim_and_pending_outbox_are_created_atomically(self):
        import sqlite3

        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["news"], frequency="once_daily"
        )
        original_connect = self.service._connect
        injected = False

        class FailingConnection:
            def __init__(self, connection):
                self.connection = connection

            def __getattr__(self, name):
                return getattr(self.connection, name)

            def execute(self, sql, parameters=()):
                nonlocal injected
                if not injected and "INSERT INTO pending_subscription_deliveries" in sql:
                    injected = True
                    raise sqlite3.OperationalError("simulated process failure before outbox")
                return self.connection.execute(sql, parameters)

            def close(self):
                self.connection.close()

        self.service._connect = lambda: FailingConnection(original_connect())
        with self.assertRaisesRegex(sqlite3.OperationalError, "simulated process failure"):
            self.service.dispatch_news_after_crawl(
                crawl_slot="2099-01-01@03:00",
                slot_label="晨间扫描",
                items=[{"title": "原子入队新闻"}],
            )
        self.service._connect = original_connect

        with sqlite3.connect(self.service.db_path) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM news_crawl_dispatches").fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM deliveries").fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM pending_subscription_deliveries").fetchone()[0], 0)

        recovered = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@03:00",
            slot_label="晨间扫描",
            items=[{"title": "原子入队新闻"}],
        )
        self.assertEqual(recovered["queued_count"], 1)

    @mock.patch.object(SubscriptionService, "_deliver_one", return_value=["om_clock_test"])
    def test_two_people_keep_independent_personal_news_times(self, delivery):
        self.service.save_subscriptions(
            "ou_delivery123", "甲", ["news"], frequency="twice_daily",
            news_delivery_times=["08:15", "18:30"],
        )
        self.service.save_subscriptions(
            "ou_second123", "乙", ["news"], frequency="twice_daily",
            news_delivery_times=["09:05", "19:10"],
        )

        queued = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-02@03:00",
            slot_label="晨间扫描",
            items=[{"title": "分时新闻"}],
            completed_at="2099-01-02T06:00:00+08:00",
        )

        self.assertEqual(queued["queued_count"], 2)
        by_person = {item["open_id"]: item["due_at"] for item in queued["results"]}
        self.assertEqual(by_person["ou_delivery123"], "2099-01-02T08:15:00+08:00")
        self.assertEqual(by_person["ou_second123"], "2099-01-02T09:05:00+08:00")
        first = self.service.flush_due(prepared_news_only=False,
            now=datetime.fromisoformat("2099-01-02T08:15:00+08:00")
        )
        self.assertEqual(first["verified_count"], 1)
        second = self.service.flush_due(prepared_news_only=False,
            now=datetime.fromisoformat("2099-01-02T09:05:00+08:00")
        )
        self.assertEqual(second["verified_count"], 1)

    @mock.patch.object(SubscriptionService, "_deliver_one", return_value=["om_clock_test"])
    def test_twice_daily_dispatch_blocks_same_window_after_schedule_change(self, delivery):
        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["news"], frequency="twice_daily"
        )
        self.service.update_news_schedule(enabled=True)
        item = {"title": "下午茶", "summary": "已完成审核"}

        first = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@13:30",
            slot_label="午后扫描",
            items=[item],
        )
        shifted = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@14:00",
            slot_label="午后扫描",
            items=[item],
        )

        self.assertEqual(first["queued_count"], 1)
        self.assertEqual(shifted["skipped_count"], 1)
        self.service.flush_due(prepared_news_only=False, now=datetime.fromisoformat("2099-01-01T19:00:00+08:00"))
        delivery.assert_called_once()
        self.assertEqual(delivery.call_args.kwargs["title"], "CMHK战略下午茶订阅｜2099年01月01日")

    def test_twice_daily_dispatch_honors_legacy_exact_slot_claim(self):
        import sqlite3

        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["news"], frequency="twice_daily"
        )
        self.service.update_news_schedule(enabled=True)
        with sqlite3.connect(self.service.db_path) as db:
            db.execute(
                """INSERT INTO news_crawl_dispatches(
                       open_id, dispatch_key, crawl_slot, crawl_date, frequency,
                       status, created_at, updated_at
                   ) VALUES(?, ?, ?, ?, ?, 'verified', ?, ?)""",
                (
                    "ou_delivery123", "twice_daily:2099-01-01@13:30",
                    "2099-01-01@13:30", "2099-01-01", "twice_daily",
                    "2099-01-01T13:31:00+08:00", "2099-01-01T13:31:00+08:00",
                ),
            )

        result = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@14:00",
            slot_label="午后扫描",
            items=[{"title": "不应重发"}],
        )

        self.assertEqual(result["skipped_count"], 1)
        self.assertFalse(any("+messages-send" in call for call in self.lark.calls))

    def test_news_dispatch_sorts_globally_and_applies_subscriber_item_limit(self):
        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["news"],
            frequency="twice_daily", news_item_limit=5,
        )
        self.service.update_news_schedule(enabled=True)
        items = [
            {
                "title": f"新闻{hour}",
                "category": "公司动态",
                "published_at": f"2099-01-03T{hour:02d}:00:00+08:00",
                "source_url": f"https://example.test/{hour}",
            }
            for hour in (3, 8, 1, 7, 2, 6, 4, 5)
        ]
        with mock.patch.object(self.service, "_deliver_one", return_value=["om_test123"]) as deliver:
            result = self.service.dispatch_news_after_crawl(
                crawl_slot="2099-01-03@09:00",
                slot_label="晨间扫描",
                items=items,
            )
            self.service.flush_due(prepared_news_only=False, now=datetime.fromisoformat("2099-01-03T08:00:00+08:00"))

        body = deliver.call_args.kwargs["body"]
        delivered = json.loads(body.removeprefix("CMHK_NEWS_DIGEST_V1\n"))
        self.assertEqual([item["title"] for item in delivered], ["新闻8", "新闻7", "新闻6", "新闻5", "新闻4"])
        self.assertEqual(deliver.call_args.kwargs["title"], "CMHK战略早茶订阅｜2099年01月03日")
        self.assertEqual(result["results"][0]["news_item_limit"], 5)

    def test_news_dispatch_filters_each_recipient_then_groups_matching_categories(self):
        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["news"],
            frequency="twice_daily", news_item_limit=5,
            news_categories=["竞对动态", "政策监管"],
        )
        self.service.update_news_schedule(enabled=True)
        items = [
            {"title": "最新行业", "category": "行业动态", "published_at": "2099-01-04T12:00:00+08:00"},
            {"title": "最新竞对", "category": "竞对动态", "published_at": "2099-01-04T11:00:00+08:00"},
            {"title": "最新政策", "category": "政策监管", "published_at": "2099-01-04T10:00:00+08:00"},
            {"title": "旧竞对", "category": "竞对动态", "published_at": "2099-01-04T09:00:00+08:00"},
        ]
        with mock.patch.object(self.service, "_deliver_one", return_value=["om_test123"]) as deliver:
            result = self.service.dispatch_news_after_crawl(
                crawl_slot="2099-01-04@14:00",
                slot_label="午后扫描",
                items=items,
            )
            self.service.flush_due(prepared_news_only=False, now=datetime.fromisoformat("2099-01-04T18:30:00+08:00"))

        delivered = json.loads(deliver.call_args.kwargs["body"].removeprefix("CMHK_NEWS_DIGEST_V1\n"))
        self.assertEqual([item["title"] for item in delivered], ["最新竞对", "最新政策", "旧竞对"])
        self.assertEqual(deliver.call_args.kwargs["title"], "CMHK战略下午茶订阅｜2099年01月04日")
        self.assertEqual(result["results"][0]["news_categories"], ["竞对动态", "政策监管"])

    def test_afternoon_news_never_repeats_for_recipient_and_backfills_from_morning_pool(self):
        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["news"],
            frequency="twice_daily", news_item_limit=5,
            news_categories=["公司动态"],
        )
        self.service.update_news_schedule(enabled=True)
        morning_items = [
            {
                "news_id": f"morning-{index}",
                "title": f"上午新闻{index}",
                "category": "公司动态",
                "published_at": f"2099-01-05T{12-index:02d}:00:00+08:00",
                "source_url": f"https://example.test/morning/{index}",
            }
            for index in range(1, 7)
        ]
        self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-05@03:00",
            slot_label="晨间扫描",
            items=morning_items,
        )
        afternoon = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-05@14:00",
            slot_label="午后扫描",
            items=[
                {**morning_items[0], "source_url": morning_items[0]["source_url"] + "?utm_source=repeat#top"},
                {
                    "news_id": "afternoon-new",
                    "title": "下午新增",
                    "category": "公司动态",
                    "published_at": "2099-01-05T14:00:00+08:00",
                    "source_url": "https://example.test/afternoon/new",
                },
            ],
        )

        self.assertEqual(afternoon["queued_count"], 1)
        with sqlite3.connect(self.service.db_path) as db:
            rows = db.execute(
                """SELECT body FROM pending_subscription_deliveries
                   WHERE open_id='ou_delivery123' ORDER BY id"""
            ).fetchall()
        morning_sent = json.loads(rows[0][0].removeprefix(NEWS_DIGEST_PREFIX))
        afternoon_sent = json.loads(rows[1][0].removeprefix(NEWS_DIGEST_PREFIX))
        morning_ids = {item["news_id"] for item in morning_sent}
        afternoon_ids = {item["news_id"] for item in afternoon_sent}
        self.assertTrue(morning_ids.isdisjoint(afternoon_ids))
        self.assertEqual([item["news_id"] for item in afternoon_sent], ["afternoon-new", "morning-6"])

    def test_daily_news_history_is_isolated_per_recipient(self):
        self.service.save_subscriptions(
            "ou_delivery123", "甲", ["news"], frequency="twice_daily",
            news_item_limit=5, news_categories=["公司动态"],
        )
        self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-06@03:00",
            slot_label="晨间扫描",
            items=[{"news_id": "shared", "title": "共同新闻", "category": "公司动态", "published_at": "2099-01-06T08:00:00+08:00"}],
        )
        self.service.save_subscriptions(
            "ou_second123", "乙", ["news"], frequency="twice_daily",
            news_item_limit=5, news_categories=["公司动态"],
        )
        self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-06@14:00",
            slot_label="午后扫描",
            items=[
                {"news_id": "shared", "title": "共同新闻", "category": "公司动态", "published_at": "2099-01-06T08:00:00+08:00"},
                {"news_id": "new", "title": "新增新闻", "category": "公司动态", "published_at": "2099-01-06T08:00:00+08:00"},
            ],
        )

        with sqlite3.connect(self.service.db_path) as db:
            rows = db.execute(
                """SELECT open_id, body FROM pending_subscription_deliveries
                   WHERE content_ref='strategic-crawl:2099-01-06@14:00' ORDER BY open_id"""
            ).fetchall()
        delivered = {
            open_id: [item["news_id"] for item in json.loads(body.removeprefix(NEWS_DIGEST_PREFIX))]
            for open_id, body in rows
        }
        self.assertEqual(delivered["ou_delivery123"], ["new"])
        self.assertEqual(delivered["ou_second123"], ["shared", "new"])

    def test_new_subscribers_default_to_four_news_categories(self):
        self.service.save_subscriptions("ou_delivery123", "测试用户", ["news"])
        summary = self.service.list_summary()
        self.assertEqual(len(summary["subscribers"][0]["news_categories"]), 4)
        self.assertEqual(len(summary["news_categories"]), 7)

    def test_invalid_preferences_are_adjusted_and_saved(self):
        saved = self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["news"], news_categories=[],
            frequency="hourly", report_mode="voice_note", news_item_limit=12,
            news_delivery_times=["07:30", "13:30"])
        self.assertEqual(saved["frequency"], "once_daily")
        self.assertEqual(saved["report_mode"], "pdf")
        self.assertEqual(saved["news_item_limit"], 10)
        self.assertEqual(len(saved["news_categories"]), 4)
        self.assertEqual(saved["news_delivery_times"], ["08:00", "14:00"])
        self.assertEqual(len(saved["adjustments"]), 6)
        actual = self.service.list_summary()["subscribers"][0]
        self.assertEqual(actual["news_delivery_times"], saved["news_delivery_times"])

    def test_overselected_callback_saves_and_sends_receipt_with_consequences(self):
        self.service.publish_entry_card(target_id="oc_test123", target_type="chat")
        event = {"type": "card.action.trigger", "action_tag": "button", "event_id": "correction-1",
            "operator_id": "ou_callback123", "chat_id": "oc_test123", "message_id": "om_test123",
            "form_value": json.dumps({"services": ["news", "weekly"], "news_categories": list(NEWS_CATEGORY_LABELS),
                "news_frequency": "twice_daily", "news_item_limit": "20", "report_mode": "pdf_audio",
                "news_delivery_time_morning": "07:59 +0800", "news_delivery_time_afternoon": "13:59 +0800"})}
        saved = self.service.handle_card_event(event)
        self.assertEqual(saved["status"], "subscription_saved")
        self.assertEqual(len(saved["news_categories"]), 7)
        self.assertIn("竞对动态", saved["news_categories"])
        self.assertEqual(saved["news_delivery_times"], ["08:00", "14:00"])
        sends = [c for c in self.lark.calls if "+messages-send" in c]
        card = json.loads(sends[-1][sends[-1].index("--content")+1])
        self.assertIn("已自动调整并保存", json.dumps(card, ensure_ascii=False))
        self.assertIn("从有新内容的已选板块中挑选最多4个", json.dumps(card, ensure_ascii=False))
        for category in saved["news_category_labels"]:
            self.assertIn(category, json.dumps(card, ensure_ascii=False))
        again = self.service.handle_card_event(event)
        self.assertEqual(again["news_categories"], saved["news_categories"])

    def test_rejected_group_submission_is_visible_and_a_retry_clears_the_error(self):
        self.service.publish_entry_card(target_id="oc_test123", target_type="chat")
        event = {
            "type": "card.action.trigger",
            "action_tag": "button",
            "event_id": "group-needs-correction",
            "operator_id": "ou_callback123",
            "chat_id": "oc_test123",
            "message_id": "om_test123",
            "form_value": json.dumps({"services": []}),
        }
        with self.assertRaisesRegex(ValueError, "至少选择") as caught:
            self.service.handle_card_event(event)
        feedback = self.service.subscription_validation_feedback(event, caught.exception)
        self.assertEqual(feedback["status"], "subscription_rejected")
        self.assertEqual(feedback["display_name"], "测试用户")
        summary = self.service.list_summary()
        group = summary["group_invitations"][0]
        self.assertEqual(group["response_count"], 1)
        self.assertEqual(group["accepted_count"], 0)
        self.assertEqual(group["responses"][0]["status"], "needs_correction")
        self.assertIn("至少选择", group["responses"][0]["last_error"])

        event["event_id"] = "group-corrected"
        event["form_value"] = json.dumps({
            "services": ["news"],
            "news_categories": list(NEWS_CATEGORY_LABELS),
        })
        saved = self.service.handle_card_event(event)
        self.assertEqual(len(saved["news_categories"]), 7)
        group = self.service.list_summary()["group_invitations"][0]
        self.assertEqual(group["accepted_count"], 1)
        self.assertEqual(group["responses"][0]["status"], "accepted")
        self.assertEqual(group["responses"][0]["last_error"], "")

    def test_invalid_times_use_defaults_and_valid_boundaries_stay(self):
        for supplied, expected in [(["bad", ""], ["08:00", "18:30"]), (["20:00", "15:00"], ["08:00", "18:30"]), (["08:00", "14:00"], ["08:00", "14:00"])]:
            with self.subTest(supplied=supplied):
                saved = self.service.save_subscriptions("ou_delivery123", "测试用户", ["news"], news_delivery_times=supplied)
                self.assertEqual(saved["news_delivery_times"], expected)

    def test_weekly_report_preference_persists_and_is_exposed_in_summary(self):
        saved = self.service.update_weekly_report_preference("reports/战略周报（编辑稿）.docx")
        self.assertEqual(saved["path"], "reports/战略周报（编辑稿）.docx")

        reloaded = SubscriptionService(runtime_root=self.root, command_runner=self.lark)
        self.assertEqual(
            reloaded.weekly_report_preference()["path"],
            "reports/战略周报（编辑稿）.docx",
        )
        self.assertEqual(
            reloaded.list_summary()["weekly_report_preference"]["path"],
            "reports/战略周报（编辑稿）.docx",
        )
        with self.assertRaisesRegex(ValueError, "路径无效"):
            self.service.update_weekly_report_preference("../secret.docx")

    def test_performance_report_preference_persists_and_is_exposed_in_summary(self):
        saved = self.service.update_performance_report_preference("reports/运营商业绩摘要（编辑稿）.docx")
        self.assertEqual(saved["path"], "reports/运营商业绩摘要（编辑稿）.docx")

        reloaded = SubscriptionService(runtime_root=self.root, command_runner=self.lark)
        self.assertEqual(
            reloaded.performance_report_preference()["path"],
            "reports/运营商业绩摘要（编辑稿）.docx",
        )
        self.assertEqual(
            reloaded.list_summary()["performance_report_preference"]["path"],
            "reports/运营商业绩摘要（编辑稿）.docx",
        )
        with self.assertRaisesRegex(ValueError, "路径无效"):
            self.service.update_performance_report_preference("../secret.docx")

    def test_report_schedule_persists_multiple_month_days_and_hong_kong_time(self):
        from datetime import datetime

        initial = self.service.report_schedule_snapshot(
            now=datetime.fromisoformat("2026-08-19T08:00:00+08:00")
        )
        self.assertFalse(initial["enabled"])
        self.assertEqual(initial["days"], [15, 30])
        saved = self.service.update_report_schedule(days="30，15, 15", time_hm="09:30", enabled=True)
        self.assertEqual(saved["days"], [15, 30])
        self.assertEqual(saved["time"], "09:30")
        self.assertTrue(saved["enabled"])
        snapshot = self.service.report_schedule_snapshot(
            now=datetime.fromisoformat("2026-08-19T08:00:00+08:00")
        )
        self.assertEqual(snapshot["next_run_at"], "2026-08-30T09:30+08:00")

    def test_report_schedule_is_due_once_after_configured_time(self):
        from datetime import datetime

        self.service.update_report_schedule(days=[15, 30], time_hm="09:30", enabled=True)
        before = self.service.report_schedule_due(
            now=datetime.fromisoformat("2026-08-30T09:29:00+08:00")
        )
        due = self.service.report_schedule_due(
            now=datetime.fromisoformat("2026-08-30T09:30:00+08:00")
        )
        self.assertFalse(before["due"])
        self.assertTrue(due["due"])
        self.assertEqual(due["slot"], "2026-08-30@09:30")

    def test_report_schedule_rejects_invalid_days_and_time(self):
        with self.assertRaisesRegex(ValueError, "1 至 31"):
            self.service.update_report_schedule(days="0, 15", time_hm="09:00", enabled=True)
        with self.assertRaisesRegex(ValueError, "HH:MM"):
            self.service.update_report_schedule(days="15, 30", time_hm="25:00", enabled=True)

    def test_due_performance_schedule_delivers_selected_version_only_once(self):
        from datetime import datetime

        report = self.root / "8月30日运营商业绩摘要（编辑稿）.docx"
        report.write_bytes(b"selected performance")
        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["performance"], report_mode="pdf"
        )
        self.service.update_performance_report_preference(report.name)
        self.service.update_performance_schedule(days=[30], time_hm="09:30", enabled=True)
        now = datetime.fromisoformat("2026-08-30T09:30:00+08:00")
        delivery = {
            "recipient_count": 1,
            "verified_count": 1,
            "queued_count": 0,
            "failed_count": 0,
        }
        with mock.patch.object(self.service, "push", return_value=delivery) as push:
            first = self.service.run_due_performance_report(now=now)
            second = self.service.run_due_performance_report(now=now)

        self.assertTrue(first["ok"])
        self.assertEqual(first["status"], "verified")
        self.assertEqual(first["selection"], "manual")
        self.assertEqual(first["report_path"], report.name)
        self.assertFalse(second["due"])
        push.assert_called_once()
        self.assertEqual(push.call_args.kwargs["service"], "performance")
        self.assertEqual(push.call_args.kwargs["path"], report.name)

    def test_due_report_schedule_generates_and_delivers_only_once_per_slot(self):
        from datetime import datetime

        self.service.save_subscriptions(
            "ou_delivery123", "测试用户", ["weekly"], report_mode="pdf"
        )
        self.service.update_report_schedule(days=[30], time_hm="09:30", enabled=True)

        def run_with_generated_report(argv, timeout=45):
            if any("generate_weekly_report.py" in item for item in argv):
                report_path = self.root / "8月30日周报.docx"
                self._write_weekly_report(report_path)
                self._write_weekly_quality_sidecar(report_path)
                return subprocess.CompletedProcess(argv, 0, "generated", "")
            return self.lark(argv, timeout=timeout)

        self.service.command_runner = run_with_generated_report
        now = datetime.fromisoformat("2026-08-30T09:30:00+08:00")
        with mock.patch.object(self.service, "_deliver_one", return_value=["om_scheduled123"]):
            first = self.service.run_due_weekly_report(now=now)
            second = self.service.run_due_weekly_report(now=now)
        self.assertTrue(first["ok"])
        self.assertEqual(first["status"], "verified")
        self.assertEqual(first["report_path"], "8月30日周报.docx")
        self.assertEqual(first["quality_gate"]["included"], 4)
        self.assertEqual(first["quality_gate"]["min_detail_sentences"], 2)
        self.assertFalse(second["due"])
        self.assertEqual(self.service.report_schedule_snapshot(now=now)["last_slot"], "2026-08-30@09:30")

    def test_due_report_selects_the_newest_generated_quality_gated_version(self):
        from datetime import datetime

        self.service.save_subscriptions("ou_delivery123", "测试用户", ["weekly"], report_mode="pdf")
        self.service.update_report_schedule(days=[30], time_hm="09:30", enabled=True)

        def run_with_two_generated_reports(argv, timeout=45):
            if any("generate_weekly_report.py" in item for item in argv):
                first = self.root / "8月30日周报.docx"
                latest = self.root / "8月30日周报 (1).docx"
                self._write_weekly_report(first)
                self._write_weekly_report(
                    latest,
                    detail=(
                        "最新报告披露经审核的业务进展，并列明参与范围、执行时点与当前状态。"
                        "后续说明补充可核验数字、具体安排和业务影响，供正式订阅者阅读。"
                    ),
                )
                self._write_weekly_quality_sidecar(first)
                self._write_weekly_quality_sidecar(latest, detail_chars=110)
                first_mtime = first.stat().st_mtime_ns
                os.utime(latest, ns=(first_mtime + 1_000_000, first_mtime + 1_000_000))
                return subprocess.CompletedProcess(argv, 0, "generated", "")
            return self.lark(argv, timeout=timeout)

        self.service.command_runner = run_with_two_generated_reports
        with mock.patch.object(self.service, "_deliver_one", return_value=["om_latest123"]):
            result = self.service.run_due_weekly_report(
                now=datetime.fromisoformat("2026-08-30T09:30:00+08:00")
            )
        self.assertTrue(result["ok"])
        self.assertEqual(result["report_path"], "8月30日周报 (1).docx")
        self.assertEqual(result["quality_gate"]["min_detail_chars"], 110)

    def test_due_report_rejects_generated_word_without_a_bound_quality_audit(self):
        from datetime import datetime

        self.service.save_subscriptions("ou_delivery123", "测试用户", ["weekly"], report_mode="pdf")
        self.service.update_report_schedule(days=[30], time_hm="09:30", enabled=True)

        def run_without_sidecar(argv, timeout=45):
            if any("generate_weekly_report.py" in item for item in argv):
                (self.root / "8月30日周报.docx").write_bytes(b"unreviewed")
                return subprocess.CompletedProcess(argv, 0, "generated", "")
            return self.lark(argv, timeout=timeout)

        self.service.command_runner = run_without_sidecar
        with mock.patch.object(self.service, "_deliver_one") as deliver:
            result = self.service.run_due_weekly_report(
                now=datetime.fromisoformat("2026-08-30T09:30:00+08:00")
            )
        self.assertFalse(result["ok"])
        self.assertIn("缺少同名质量审计", result["error"])
        deliver.assert_not_called()

    def test_due_report_rejects_legacy_one_sentence_audit_even_if_marked_passed(self):
        from datetime import datetime

        self.service.save_subscriptions("ou_delivery123", "测试用户", ["weekly"], report_mode="pdf")
        self.service.update_report_schedule(days=[30], time_hm="09:30", enabled=True)

        def run_with_legacy_sidecar(argv, timeout=45):
            if any("generate_weekly_report.py" in item for item in argv):
                report_path = self.root / "8月30日周报.docx"
                report_path.write_bytes(b"legacy")
                self._write_weekly_quality_sidecar(
                    report_path,
                    detail_chars=20,
                    detail_sentences=None,
                )
                return subprocess.CompletedProcess(argv, 0, "generated", "")
            return self.lark(argv, timeout=timeout)

        self.service.command_runner = run_with_legacy_sidecar
        with mock.patch.object(self.service, "_deliver_one") as deliver:
            result = self.service.run_due_weekly_report(
                now=datetime.fromisoformat("2026-08-30T09:30:00+08:00")
            )
        self.assertFalse(result["ok"])
        self.assertIn("至少 90 字且 2 句完整事实", result["error"])
        deliver.assert_not_called()

    def test_due_report_rejects_quality_audit_bound_to_different_word_bytes(self):
        from datetime import datetime

        self.service.save_subscriptions("ou_delivery123", "测试用户", ["weekly"], report_mode="pdf")
        self.service.update_report_schedule(days=[30], time_hm="09:30", enabled=True)

        def run_with_stale_sidecar(argv, timeout=45):
            if any("generate_weekly_report.py" in item for item in argv):
                report_path = self.root / "8月30日周报.docx"
                report_path.write_bytes(b"latest")
                self._write_weekly_quality_sidecar(report_path, report_sha256="0" * 64)
                return subprocess.CompletedProcess(argv, 0, "generated", "")
            return self.lark(argv, timeout=timeout)

        self.service.command_runner = run_with_stale_sidecar
        with mock.patch.object(self.service, "_deliver_one") as deliver:
            result = self.service.run_due_weekly_report(
                now=datetime.fromisoformat("2026-08-30T09:30:00+08:00")
            )
        self.assertFalse(result["ok"])
        self.assertIn("哈希与待推送 Word 不一致", result["error"])
        deliver.assert_not_called()

    def test_weekly_push_rejects_navigation_noise_even_with_passed_sidecar(self):
        report = self.root / "含网页导航噪声周报.docx"
        self._write_weekly_report(
            report,
            detail=(
                "发掘 SmarTone 精选服务计划，5G计划任你拣，总有一个适合你。"
                "数据用量 - 高至低。数据用量 - 低至高。价格 - 高至低。价格 - 低至高。"
            ),
        )
        self._write_weekly_quality_sidecar(report)

        with self.assertRaisesRegex(RuntimeError, "排序、导航或营销按钮噪声"):
            self.service.push(
                service="weekly",
                mode="pdf",
                path=report.name,
                test_open_id="ou_test123",
            )
        self.assertFalse(any("+messages-send" in call for call in self.lark.calls))

    def test_due_report_does_not_generate_without_an_active_weekly_subscriber(self):
        from datetime import datetime

        self.service.update_report_schedule(days=[30], time_hm="09:30", enabled=True)
        now = datetime.fromisoformat("2026-08-30T09:30:00+08:00")
        result = self.service.run_due_weekly_report(now=now)
        self.assertEqual(result["skipped"], "no_active_subscribers")
        self.assertFalse(any("generate_weekly_report.py" in item for call in self.lark.calls for item in call))

    def test_due_report_never_falls_back_to_an_old_report(self):
        from datetime import datetime

        (self.root / "旧周报.docx").write_bytes(b"old")
        self.service.save_subscriptions("ou_delivery123", "测试用户", ["weekly"], report_mode="pdf")
        self.service.update_report_schedule(days=[30], time_hm="09:30", enabled=True)

        def successful_but_no_output(argv, timeout=45):
            if any("generate_weekly_report.py" in item for item in argv):
                return subprocess.CompletedProcess(argv, 0, "generated", "")
            return self.lark(argv, timeout=timeout)

        self.service.command_runner = successful_but_no_output
        result = self.service.run_due_weekly_report(
            now=datetime.fromisoformat("2026-08-30T09:30:00+08:00")
        )
        self.assertFalse(result["ok"])
        self.assertIn("本轮新生成", result["error"])
        self.assertFalse(any("+messages-send" in call for call in self.lark.calls))

    def test_news_dispatch_requires_active_subscription(self):
        self.service.save_subscriptions("ou_delivery123", "测试用户", ["news"], frequency="once_daily")
        queued = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@07:00", slot_label="晨间扫描", items=[{"title": "待发送"}]
        )
        self.assertEqual(queued["queued_count"], 1)
        self.service.update_subscriber(
            "ou_delivery123", services=["news"], status="paused", frequency="once_daily", report_mode="pdf"
        )
        unsubscribed = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@15:00", slot_label="午后扫描", items=[{"title": "仍不发送"}]
        )
        self.assertTrue(unsubscribed["schedule_enabled"])
        self.assertEqual(unsubscribed["recipient_count"], 0)

    @mock.patch.object(SubscriptionService, "_deliver_one", side_effect=RuntimeError("temporary offline"))
    def test_failed_crawler_delivery_retries_without_a_fixed_attempt_cap(self, delivery):
        from datetime import datetime
        import sqlite3

        self.service.save_subscriptions("ou_delivery123", "测试用户", ["news"], frequency="once_daily")
        self.service.update_news_schedule(enabled=True)

        def offline(argv, timeout=45):
            raise RuntimeError("temporary offline")

        self.service.command_runner = offline
        queued = self.service.dispatch_news_after_crawl(
            crawl_slot="2099-01-01@07:00",
            slot_label="晨间扫描",
            items=[{"title": "重试新闻"}],
        )
        self.assertEqual(queued["queued_count"], 1)
        first = self.service.flush_due(prepared_news_only=False, now=datetime.fromisoformat("2099-01-01T19:00:00+08:00"))
        second = self.service.flush_due(prepared_news_only=False, now=datetime.fromisoformat("2099-01-01T20:00:00+08:00"))
        self.assertEqual(first["retrying_count"], 1)
        self.assertEqual(second["retrying_count"], 1)
        db = sqlite3.connect(self.service.db_path)
        try:
            status, attempts = db.execute(
                "SELECT status, attempts FROM pending_subscription_deliveries"
            ).fetchone()
            delivery_status = db.execute("SELECT status FROM deliveries").fetchone()[0]
        finally:
            db.close()
        self.assertEqual(status, "queued")
        self.assertEqual(attempts, 2)
        self.assertEqual(delivery_status, "retrying")

    def test_long_report_paragraphs_are_split_below_message_limit(self):
        chunks = self.service._text_chunks("长报告", "甲" * 12000)
        self.assertGreater(len(chunks), 2)
        self.assertTrue(all(len(chunk) <= 5800 for chunk in chunks))


if __name__ == "__main__":
    unittest.main()
