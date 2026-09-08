import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

import web_app
from cmhk.services.subscriptions import SubscriptionService


class NewReportSelectionTests(unittest.TestCase):
    def test_files_added_after_first_listing_are_selectable_and_persist_across_reopen(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            service = SubscriptionService(runtime_root=root)

            # Keep the real disk discovery and report metadata mapping; omit
            # unrelated crawler/dashboard metrics from this focused test.
            def report_status():
                return {"outputs": [web_app.file_info(p) for p in web_app.current_report_files()]}

            with (
                mock.patch.object(web_app, "ROOT", root),
                mock.patch.object(web_app, "REPORT_METADATA_PATH", root / "data/reporting/report_file_metadata.json"),
                mock.patch.object(web_app, "report_audio_metadata", return_value={"exists": False}),
                mock.patch.object(web_app, "build_status", side_effect=report_status),
            ):
                self.assertEqual(report_status()["outputs"], [])
                for wave in (1, 2):
                    for report_type, title, update, read in (
                        ("weekly", "未来新增周报", web_app.update_weekly_report_preference, web_app.weekly_report_preference_payload),
                        ("carrier-performance", "未来新增运营商业绩摘要", web_app.update_performance_report_preference, web_app.performance_report_preference_payload),
                    ):
                        for edited in (False, True):
                            with self.subTest(wave=wave, report_type=report_type, edited=edited):
                                name = f"{title} ({wave}){'（编辑稿）' if edited else ''}.docx"
                                before = {item["path_str"] for item in report_status()["outputs"]}
                                self.assertNotIn(name, before)
                                doc = Document()
                                doc.add_paragraph(f"第{wave}轮新增的完整报告正文，保存后应立即进入报告库并允许选择。")
                                doc.save(root / name)
                                result = update(service, name)
                                self.assertEqual(result["path"], name)
                                self.assertTrue(result["available"])
                                self.assertEqual(result["report"]["is_edited"], edited)
                                reopened = SubscriptionService(runtime_root=root)
                                self.assertEqual(read(reopened)["path"], name)
                                self.assertFalse(Path(str(root / name) + ".quality.json").exists())
                self.assertEqual(len(report_status()["outputs"]), 8)

    def test_new_invalid_file_does_not_replace_the_previous_selection(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            service = SubscriptionService(runtime_root=root)
            for title, report_type, update, read in (
                ("周报", "weekly", web_app.update_weekly_report_preference, service.weekly_report_preference),
                ("运营商业绩摘要", "carrier-performance", web_app.update_performance_report_preference, service.performance_report_preference),
            ):
                with self.subTest(report_type=report_type):
                    valid = f"有效{title}.docx"
                    invalid = f"新增损坏{title}.docx"
                    doc = Document()
                    doc.add_paragraph("原先选择的有效报告。")
                    doc.save(root / valid)
                    (root / invalid).write_bytes(b"broken docx")
                    status = {"outputs": [{"path_str": p, "reportType": report_type} for p in (valid, invalid)]}
                    with mock.patch.object(web_app, "build_status", return_value=status):
                        update(service, valid)
                        with self.assertRaisesRegex(ValueError, "无法读取"):
                            update(service, invalid)
                    self.assertEqual(read()["path"], valid)


if __name__ == "__main__":
    unittest.main()
