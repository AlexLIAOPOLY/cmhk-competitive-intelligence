from __future__ import annotations

# Annotation imports do not create application services or mutable state.
from pathlib import Path

from ._binding import publish


def bind(app) -> None:
    """Bind this domain to the existing application and its shared state."""
    def reference_path(name: str) -> Path | None:
        raw = str(name or "").strip().lstrip("/")
        clean = app.Path(raw).name
        if clean in app.REFERENCE_FILES:
            return app.ROOT / clean
        if app.re.fullmatch(r"row_\d+\.json", clean):
            return app.RESULTS_DIR / clean
        if raw.startswith("agent_knowledge/"):
            target = (app.ROOT / raw).resolve()
            knowledge_root = (app.ROOT / "agent_knowledge").resolve()
            if knowledge_root in target.parents and target.exists() and target.is_file():
                return target
        return None

    publish(app, reference_path)

    def decode_text_bytes(body: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5", "cp950"):
            try:
                return body.decode(encoding)
            except UnicodeDecodeError:
                continue
        return body.decode("utf-8", errors="replace")

    publish(app, decode_text_bytes)

    def read_display_text(path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".docx":
            try:
                from docx import Document

                doc = Document(str(path))
                parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            parts.append(" | ".join(cells))
                return "\n".join(parts)
            except Exception as exc:
                return f"Word 文档预览失败：{exc}"
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader

                reader = PdfReader(str(path))
                return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
            except Exception as exc:
                return f"PDF 预览失败：{exc}"
        raw = app.decode_text_bytes(path.read_bytes())
        if suffix == ".json":
            try:
                raw = app.json.dumps(app.json.loads(raw), ensure_ascii=False, indent=2)
            except Exception:
                pass
        return raw

    publish(app, read_display_text)

    def load_report_metadata() -> dict:
        if not app.REPORT_METADATA_PATH.exists():
            return {}
        try:
            data = app.json.loads(app.REPORT_METADATA_PATH.read_text(encoding="utf-8"))
        except Exception:
            return {}
        return data if isinstance(data, dict) else {}

    publish(app, load_report_metadata)

    def save_report_metadata(data: dict) -> None:
        app.REPORT_METADATA_PATH.write_text(app.json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    publish(app, save_report_metadata)

    def is_report_path(path: Path) -> bool:
        if not path.exists() or not path.is_file() or path.suffix.lower() != ".docx":
            return False
        if path.name.startswith("~$") or path.name in app.EXCLUDED_REPORT_NAMES:
            return False
        try:
            path.relative_to(app.ROOT)
        except ValueError:
            return False
        return path.parent == app.ROOT or app.ROOT / "archives" in path.parents

    publish(app, is_report_path)

    def quality_sidecar_for_report(path: Path) -> Path:
        return app.Path(str(path) + ".quality.json")

    publish(app, quality_sidecar_for_report)

    def report_audio_metadata(report_path: Path) -> dict:
        audio_path = next(
            (path for path in app.audio_paths_for_report(report_path) if path.exists()),
            None,
        )
        if not audio_path:
            return {"exists": False}
        audio_stat = audio_path.stat()
        return {
            "exists": True,
            "url": f"/audio/{app.quote(audio_path.name)}?v={audio_stat.st_mtime_ns}",
        }

    publish(app, report_audio_metadata)

    def file_info(path: Path, url: str = None) -> dict:
        stat = path.stat()
        rel_path = str(path.relative_to(app.ROOT))
        metadata = app.load_report_metadata().get(rel_path, {})
        metadata = metadata if isinstance(metadata, dict) else {}
        compact_audio = app.report_audio_metadata(path)
        report_type = str(metadata.get("reportType") or "")
        if report_type not in {"weekly", "carrier-performance"}:
            report_type = "carrier-performance" if "业绩摘要" in path.name else "weekly"
        return {
            "name": app.report_display_name(path.name),
            "size": stat.st_size,
            "mtime": stat.st_mtime,
            "mtimeText": app.time.strftime("%Y-%m-%d %H:%M:%S", app.time.localtime(stat.st_mtime)),
            "url": url or f"/outputs/{app.quote(path.name)}",
            "path_str": rel_path,
            "note": str(metadata.get("note") or ""),
            "reportType": report_type,
            "isEdited": bool(metadata.get("isEdited")) or "编辑稿" in path.stem,
            "editRevision": int(metadata.get("editorRevision") or 0),
            "editedAt": str(metadata.get("editedAt") or ""),
            "editedBy": str(metadata.get("editedBy") or ""),
            "sourcePath": str(metadata.get("sourcePath") or ""),
            # Full subtitle cues and spoken text are loaded only when the user
            # plays one report.  Embedding every historical transcript made the
            # ten-second status poll grow to megabytes.
            "audio": compact_audio,
        }

    publish(app, file_info)

    def is_report_file_name(name: str) -> bool:
        return name.endswith(".docx") and "/" not in name and "\\" not in name and name not in app.EXCLUDED_REPORT_NAMES

    publish(app, is_report_file_name)

    def current_report_files() -> list[Path]:
        files = [path for path in app.ROOT.glob("*.docx") if app.is_report_path(path)]
        return sorted(files, key=lambda p: p.stat().st_mtime, reverse=True)

    publish(app, current_report_files)

    def report_target_from_rel(path_str: str) -> Path | None:
        if not path_str or path_str.startswith("/") or ".." in app.Path(path_str).parts:
            return None
        target = app.ROOT / path_str
        try:
            target.relative_to(app.ROOT)
        except ValueError:
            return None
        return target if app.is_report_path(target) else None

    publish(app, report_target_from_rel)

    def update_report_file(payload: dict) -> dict:
        target = app.report_target_from_rel(str(payload.get("path") or ""))
        if not target:
            raise ValueError("文件不存在或不允许修改")
        new_name = app.Path(str(payload.get("name") or "").strip()).name
        if not new_name:
            raise ValueError("文件名不能为空")
        if not new_name.endswith(".docx"):
            new_name += ".docx"
        if not app.is_report_file_name(new_name):
            raise ValueError("文件名只能是 Word 文档，不能包含路径字符")
        if new_name == app.report_display_name(target.name):
            new_name = target.name
        new_note = app.re.sub(r"\s+", " ", str(payload.get("note") or "")).strip()[:500]
        new_target = target.with_name(new_name)
        from cmhk.reporting.report_naming import rename_report_bundle
        rename_report_bundle(app.ROOT, target, new_target, note=new_note)
        return app.build_status()

    publish(app, update_report_file)

    def _report_type_for_path(path: Path, metadata: dict | None = None) -> str:
        metadata = metadata if isinstance(metadata, dict) else {}
        saved = str(metadata.get("reportType") or "")
        if saved in {"weekly", "carrier-performance"}:
            return saved
        return "carrier-performance" if "业绩摘要" in path.name else "weekly"

    publish(app, _report_type_for_path)

    def _next_edited_report_path(source: Path) -> Path:
        base = app.re.sub(r"（编辑稿(?:\s+\d+)?）$", "", source.stem).strip()
        first = source.with_name(f"{base}（编辑稿）.docx")
        if not first.exists():
            return first
        for revision in range(2, 1000):
            candidate = source.with_name(f"{base}（编辑稿 {revision}）.docx")
            if not candidate.exists():
                return candidate
        raise ValueError("编辑稿版本过多，请先整理报告库")

    publish(app, _next_edited_report_path)

    def load_report_editor_payload(path_str: str) -> dict:
        target = app.report_target_from_rel(path_str)
        if not target:
            raise FileNotFoundError("报告不存在或不允许编辑")
        payload = app.load_docx_for_editor(target)
        rel_path = str(target.relative_to(app.ROOT))
        metadata = app.load_report_metadata().get(rel_path, {})
        metadata = metadata if isinstance(metadata, dict) else {}
        preview_url = ""
        try:
            from cmhk.reporting.pdf_preview import pdf_preview_path, convert_docx_to_pdf_preview

            preview_dir = app.ROOT / "web" / "static" / "report-previews"
            preview_path = pdf_preview_path(target, preview_dir)
            if not preview_path.is_file() or preview_path.stat().st_mtime_ns < target.stat().st_mtime_ns:
                preview_path = convert_docx_to_pdf_preview(target, preview_dir=preview_dir)
            if preview_path.is_file():
                preview_url = f"/static/report-previews/{app.quote(preview_path.name)}?v={preview_path.stat().st_mtime_ns}"
        except Exception:
            preview_url = ""
        return {
            **payload,
            "name": app.report_display_name(target.name),
            "path": rel_path,
            "reportType": app._report_type_for_path(target, metadata),
            "isEdited": bool(metadata.get("isEdited")),
            "editRevision": int(metadata.get("editorRevision") or 0),
            "editedAt": str(metadata.get("editedAt") or ""),
            "editedBy": str(metadata.get("editedBy") or ""),
            "sourcePath": str(metadata.get("sourcePath") or rel_path),
            "previewUrl": preview_url,
        }

    publish(app, load_report_editor_payload)

    def save_report_editor_payload(payload: dict, *, actor: dict | None = None) -> dict:
        source = app.report_target_from_rel(str(payload.get("path") or ""))
        if not source:
            raise FileNotFoundError("报告不存在或不允许编辑")
        expected_hash = str(payload.get("sourceSha256") or "").lower()
        if not app.re.fullmatch(r"[0-9a-f]{64}", expected_hash):
            raise ValueError("缺少编辑源版本，请重新打开报告")
        document_payload = payload.get("document")
        save_mode = str(payload.get("saveMode") or "update")
        if save_mode not in {"update", "copy"}:
            raise ValueError("不支持的保存方式")

        with app.REPORT_EDITOR_LOCK:
            current_hash = app.sha256_file(source)
            if current_hash != expected_hash:
                raise app.ReportEditConflict("这份报告在编辑期间已被其他人更新，请重新打开后再编辑")
            metadata = app.load_report_metadata()
            source_rel = str(source.relative_to(app.ROOT))
            source_meta = metadata.get(source_rel, {})
            source_meta = source_meta if isinstance(source_meta, dict) else {}
            source_is_edited = bool(source_meta.get("isEdited"))
            target = source if payload.get("bodyOnly") is True or (source_is_edited and save_mode == "update") else app._next_edited_report_path(source)

            prior_revision = int(source_meta.get("editorRevision") or 0)
            if target == source:
                history_dir = app.ROOT / "archives" / "report_edits" / app.re.sub(r"[^0-9A-Za-z\u4e00-\u9fff._-]+", "_", source.stem)
                history_dir.mkdir(parents=True, exist_ok=True)
                history_name = f"{app.datetime.now().strftime('%Y%m%d-%H%M%S')}-r{max(1, prior_revision)}.docx"
                app.shutil.copy2(source, history_dir / history_name)

            saved = app.save_editor_document(source, target, document_payload, body_only=payload.get("bodyOnly") is True)
            app.delete_audio_for_report(target)
            preview_url = ""
            warning = ""
            # Preview conversion must not open any desktop application.
            from cmhk.reporting.pdf_preview import pdf_preview_path, convert_docx_to_pdf_preview
            preview_dir = app.ROOT / "web" / "static" / "report-previews"
            stale_preview = pdf_preview_path(target, preview_dir)
            stale_preview.unlink(missing_ok=True)
            try:
                preview = convert_docx_to_pdf_preview(target, preview_dir=preview_dir)
                preview_url = f"/static/report-previews/{app.quote(preview.name)}?v={preview.stat().st_mtime_ns}"
            except Exception:
                warning = "正文已保存，PDF 预览生成失败，请稍后重新打开。"

            actor = actor if isinstance(actor, dict) else {}
            actor_name = str(actor.get("name") or actor.get("display_name") or actor.get("username") or "当前用户")[:120]
            target_rel = str(target.relative_to(app.ROOT))
            report_type = app._report_type_for_path(source, source_meta)
            root_source = str(source_meta.get("sourcePath") or source_rel)
            revision = prior_revision + 1 if target == source else 1
            metadata[target_rel] = {
                "note": str(source_meta.get("note") or f"页面编辑稿 · 来源 {app.Path(root_source).name}")[:500],
                "updatedAt": app.time.strftime("%Y-%m-%d %H:%M:%S"),
                "isEdited": True,
                "editorRevision": revision,
                "editedAt": app.datetime.now().astimezone().isoformat(timespec="seconds"),
                "editedBy": actor_name,
                "sourcePath": root_source,
                "sourceSha256": str(source_meta.get("sourceSha256") or current_hash),
                "reportType": report_type,
            }
            app.save_report_metadata(metadata)
            status = app.build_status()
            saved_file = next(
                (item for item in status.get("outputs", []) if item.get("path_str") == target_rel),
                app.file_info(target),
            )
            return {
                "file": saved_file,
                "status": status,
                "sourcePath": source_rel,
                "path": target_rel,
                "sourceSha256": str(saved["sha256"]),
                "sourceMtimeNs": int(saved["mtimeNs"]),
                "previewUrl": preview_url,
                "warning": warning,
            }

    publish(app, save_report_editor_payload)

    def delete_report_files(paths: list[str]) -> dict:
        metadata = app.load_report_metadata()
        deleted = 0
        for path_str in paths:
            target = app.report_target_from_rel(str(path_str))
            if not target:
                continue
            rel_path = str(target.relative_to(app.ROOT))
            target.unlink()
            quality_sidecar = app.quality_sidecar_for_report(target)
            if quality_sidecar.exists():
                quality_sidecar.unlink()
            app.delete_audio_for_report(target)
            metadata.pop(rel_path, None)
            deleted += 1
        app.save_report_metadata(metadata)
        return {"deleted": deleted, "status": app.build_status()}

    publish(app, delete_report_files)

    def report_overview() -> str:
        md_path = app.ROOT / "weekly_report.md"
        if not md_path.exists():
            return "当前还没有生成周报。你可以先点击“生成周报”，系统会按 Word 模板输出正式 Word 周报。"
        text = md_path.read_text(encoding="utf-8", errors="ignore")
        lines = [line.strip("#- 　\t ") for line in text.splitlines() if line.strip()]
        useful = [line for line in lines if line and not line.startswith("来源")][:8]
        status = app.build_status()
        intro = (
            "这里的周报是“战略内参周报”：把公开信息监测数据按模板整理成正式汇报文件，"
            "主要用于快速查看政策、行业、社会和国际资讯中的重点变化。"
        )
        if not useful:
            return f"{intro} 当前已有输出文件，最近生成时间是 {status['latestOutputText']}。"
        return f"{intro} 当前最近生成时间是 {status['latestOutputText']}。报告开头内容包括：" + "；".join(useful[:5]) + "。"

    publish(app, report_overview)

    def output_overview() -> str:
        status = app.build_status()
        outputs = status.get("outputs", [])
        if not outputs:
            return "当前还没有输出文件。点击“生成周报”后会生成正式 Word 周报。"
        names = "、".join(item["name"] for item in outputs)
        return f"当前可用输出文件有：{names}。这里仅展示正式 Word 周报，用于下载和提交。"

    publish(app, output_overview)

    def check_local_action(message: str) -> dict | None:
        return None

        status = app.build_status()
        status_intent = any(
            key in text
            for key in [
                "系统状态",
                "运行状态",
                "当前状态",
                "检查系统",
                "检查后端",
                "结果文件状态",
                "输出文件状态",
                "现在有多少文件",
                "现在有哪些文件",
            ]
        ) or text in {"状态", "检查", "现在", "文件"}
        if status_intent:
            return {
                "content": (
                    f"当前已有 {status['results']['count']} 个结果文件，"
                    f"ok {status['results']['ok']} 个，partial {status['results']['partial']} 个。"
                    f"模板文件{'存在' if status['template']['exists'] else '不存在'}，"
                    f"最近输出时间是 {status['latestOutputText']}。"
                ),
            }

        if "模板" in text or "格式" in text:
            return {
                "content": (
                    "当前生成流程会优先读取本地上传的模板，若无则使用库里的默认模板 weekly_report_template.docx，"
                    "保留封面、目录位置、页眉页脚和图片资源，只替换目录与正文段落文字。"
                ),
            }

        if "openai" in lowered or "api" in lowered or "ai" in lowered:
            return {
                "content": (
                    "这个助手已接入 OpenAI Responses API 的调用代码，并会先对本地周报、爬取结果和审计文件做 RAG 检索。"
                    "当前运行环境需要设置 OPENAI_API_KEY 后才能真正调用模型。"
                ),
            }

        return None

    publish(app, check_local_action)

